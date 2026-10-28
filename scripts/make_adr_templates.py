"""Write the Excel template of each type of architectural decision (X-1) in Templates/, from the fields of its sheet
in ATLAS (model.schema): the same layout, styles and print settings as the general template X-1.

    .venv/Scripts/python.exe scripts/make_adr_templates.py

Run it again whenever the fields of a type change (tests/test_adr_templates.py checks the files are up to date)."""
import sys
import unicodedata
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from arcadia_tool.model.adr import common_fields, type_fields  # noqa: E402
from arcadia_tool.model.schema import (ADR_AREAS, ADR_TYPES, ENUM, KINDS_BY_NAME, MULTIENUM, REF,  # noqa: E402
                                       REFS, TABLE, TEXT, Column, Field)

TEMPLATES = ROOT / "Templates"
GENERAL = TEMPLATES / "Template_X-1_Decisione_architetturale_ADR.xlsx"
FIRST_ROW = 11  # the rows above (title, guide, data of the document) are those of the general template
STATUSES = ("Proposta", "Decisa", "Superata")

# What the tool computes from the model (see model.adr.COMPUTED): in the template, as it is when proposed.
COMPUTED_HINTS = {
    "merged_srs_items": "le voci SRS dei domini da fondere (principale o contributore)",
    "current_mission": "la missione scritta ora nella scheda del dominio (SA-1)",
    "item_primary": "il dominio principale della voce, dal registro di tracciabilità (SA-3)",
    "item_contributors": "i domini contributori della voce, dal registro di tracciabilità (SA-3)",
    "function_domain": "il dominio della funzione, dal catalogo delle funzioni (SA-2)",
    "item_producers": "i produttori degli scambi che portano l'informazione (SA-5)",
}
# The choices of a column computed from the other fields (see model.adr.CHOICES).
CHOICE_HINTS = {
    "split_domain_items": "una voce SRS del dominio da scindere",
    "split_destinations": "uno dei domini risultanti scritti sopra",
}
# The guides of the general template, for the fields it has too.
FIELD_HINTS = {"context": "Il problema e perché va deciso ora",
               "consequences": "Cosa cambia nel modello, nei requisiti, nei test",
               "review_when": "Condizione che riapre la decisione, se esiste"}
EXAMPLES = {"domain": "DF07 · Odometria", "team": "T5 · Odometria", "requirement": "3.14.4.2",
            "system_function": "SF-DF07-03 · Stimare la velocità", "exchange_item": "EI-velocità_stimata",
            "functional_chain": "FC-frenatura_di_servizio", "logical_function": "LF-DF07-03-01",
            "logical_component": "LC-04", "adr": "ADR-nnn", "external_actor": "Conducente"}


def file_name(adr_type):
    """Template_X-1.nn_ADR_<tipo>.xlsx, numbered in the order of the areas (as in the window of the types)."""
    order = [t for _, types in ADR_AREAS for t in types]
    ascii_name = unicodedata.normalize("NFKD", adr_type).encode("ascii", "ignore").decode()
    words = "".join(c if c.isalnum() else " " for c in ascii_name).split()
    return f"Template_X-1.{order.index(adr_type) + 1:02d}_ADR_{'_'.join(words)}.xlsx"


def _examples(kinds):
    return " / ".join(f"es. {EXAMPLES[k]}" if k in EXAMPLES else KINDS_BY_NAME[k].label for k in kinds)


def hint(f):
    """The guide in grey italics in the cell of the field (to replace when the template is filled in)."""
    if f.name == "decision_date":
        return "gg/mm/aaaa · impostata da ATLAS quando la Design Authority decide"
    if f.computed:
        return f"Calcolato da ATLAS: {COMPUTED_HINTS.get(f.computed, 'dal modello')}"
    if f.by_tool:
        return "Scritto da ATLAS all'approvazione della Design Authority"
    if f.type == ENUM:
        return " / ".join(f.options)
    if f.type == MULTIENUM:
        return "Uno o più tra: " + ", ".join(f.options)
    if f.type == REF:
        return _examples(f.ref_kinds)
    if f.type == REFS:
        if f.sections_only:
            return "Una o più sezioni (tre numeri, es. 3.4.1)"
        return "Uno o più: " + _examples(f.ref_kinds)
    return FIELD_HINTS.get(f.name, "")


class Sheet:
    """The general template, emptied below the data of the document, with its styles to write the fields."""

    def __init__(self):
        self.book = load_workbook(GENERAL)
        ws = self.ws = self.book.active
        self.label, self.hint = copy(ws["A11"]._style), copy(ws["B11"]._style)
        self.hint_font = copy(ws["B11"].font)
        self.value = copy(ws["B7"]._style)
        self.header, self.cell = copy(ws["A17"]._style), copy(ws["A18"]._style)
        for merged in [r for r in ws.merged_cells.ranges if r.min_row >= FIRST_ROW]:
            ws.unmerge_cells(str(merged))
        ws.delete_rows(FIRST_ROW, ws.max_row)
        for row in [r for r in ws.row_dimensions if r >= FIRST_ROW]:
            del ws.row_dimensions[row]
        ws.data_validations.dataValidation = [v for v in ws.data_validations.dataValidation
                                              if all(r.min_row < FIRST_ROW for r in v.sqref.ranges)]
        self.row = FIRST_ROW

    def _style(self, cell, style):
        cell._style = copy(style)

    def heading(self, text):
        """A dark row across the sheet: the start of a section."""
        self.row += 1 if self.row > FIRST_ROW else 0
        ws, r = self.ws, self.row
        for col in "ABCDEF":
            self._style(ws[f"{col}{r}"], self.header)
        ws[f"A{r}"] = text
        ws.merge_cells(f"A{r}:F{r}")
        ws.row_dimensions[r].height = 20
        self.row += 1

    def field(self, label, text="", fixed="", options=(), tall=False):
        """A label and its value across B:F: the guide in grey italics, or a fixed value."""
        ws, r = self.ws, self.row
        self._style(ws[f"A{r}"], self.label)
        ws[f"A{r}"] = label
        for col in "BCDEF":
            self._style(ws[f"{col}{r}"], self.value if fixed else self.hint)
        ws[f"B{r}"] = fixed or text or None
        ws.merge_cells(f"B{r}:F{r}")
        ws.row_dimensions[r].height = 45 if tall else (30 if len(label) > 28 or len(text) > 90 else 22)
        if options:
            self.choose(f"B{r}", options)
        self.row += 1

    def choose(self, ref, options):
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        validation.add(ref)
        self.ws.add_data_validation(validation)

    def table(self, f, rows=5, spans=None):
        """The label of the table, its columns (spread over A:F, or as `spans` says) and empty rows to fill in."""
        ws = self.ws
        self._style(ws[f"A{self.row}"], self.label)
        ws[f"A{self.row}"] = f.label
        for col in "BCDEF":
            self._style(ws[f"{col}{self.row}"], self.label)
        ws.merge_cells(f"A{self.row}:F{self.row}")
        ws.row_dimensions[self.row].height = 22
        self.row += 1
        spans = spans or {1: ("AF",), 2: ("AC", "DF"), 3: ("AB", "CD", "EF"), 4: ("AA", "BC", "DD", "EF")}[
            len(f.columns)]
        for r in range(self.row, self.row + rows + 1):
            for (first, last), column in zip(spans, f.columns):
                for col in "ABCDEF"["ABCDEF".index(first):"ABCDEF".index(last) + 1]:
                    self._style(ws[f"{col}{r}"], self.header if r == self.row else self.cell)
                if r == self.row:
                    ws[f"{first}{r}"] = column.label + (f" ({CHOICE_HINTS[column.choices]})"
                                                        if column.choices in CHOICE_HINTS else "")
                elif column.options:
                    self.choose(f"{first}{r}", column.options)
                if first != last:
                    ws.merge_cells(f"{first}{r}:{last}{r}")
            ws.row_dimensions[r].height = 30 if r == self.row else 20
        self.row += rows + 1

    def write(self, f):
        if f.type == TABLE:
            self.table(f)
        else:
            self.field(f.label, hint(f), options=f.options if f.type == ENUM else (), tall=f.type == TEXT)

    def save(self, path):
        self.ws.print_area = f"A1:F{self.row - 1}"  # the one of the general template stops at its last row
        self.book.save(path)


def guide(adr_type):
    """«Come si compila»: where the ADR is opened in ATLAS and what the tool fills in by itself."""
    text = ("Come si compila. In ATLAS si crea con «Nuova ADR» e il tipo «" + adr_type + "»: la scheda ha gli stessi "
            "campi di questo template.")
    computed = [f.label for f in type_fields(adr_type) if f.computed or f.by_tool]
    if len(computed) == 1:
        text += (f" Il campo «{computed[0]}» è compilato da ATLAS dal modello: qui si riporta com'è al momento "
                 "della proposta.")
    elif computed:
        text += (" I campi " + ", ".join(f"«{c}»" for c in computed) + " sono compilati da ATLAS dal modello: qui "
                 "si riportano come sono al momento della proposta.")
    return text + " La data della decisione è quella in cui la Design Authority decide (la DA è sempre il decisore)."


def make(adr_type, when):
    sheet = Sheet()
    ws = sheet.ws
    ws["A1"] = f"X-1 · Decisione architetturale (ADR) · {adr_type}"
    ws["A2"] = ws["A2"].value.replace("v1.0 · 28/09/2026", f"v1.1 · {date.today():%d/%m/%Y}")
    ws["A3"] = f"Scopo. {when}"
    ws["A4"] = guide(adr_type)
    sheet.book.properties.title = f"X-1 · {adr_type} · template ARCADIA"

    adr = KINDS_BY_NAME["adr"]
    sheet.heading("Identificazione")
    sheet.field("ID e titolo", "ADR-nnn · titolo della decisione")
    sheet.field("Tipo di decisione", fixed=adr_type)
    sheet.field(adr.field("status").label, " / ".join(STATUSES), options=STATUSES)
    for name in ("superseded_by", "proposer", "decision_date"):
        sheet.write(adr.field(name))
    common = common_fields(adr_type)
    for f in common:  # e.g. the domain of a deletion, right under the type (as in the sheet)
        if f.section == "Identificazione":
            sheet.write(f)
    common = [f for f in common if f.section != "Identificazione"]
    details = type_fields(adr_type)
    if details:
        sheet.heading(f"Dettagli: {adr_type}")
        for f in details:
            sheet.write(f)
    if common:
        sheet.heading("Descrizione")
        for f in common:
            sheet.write(f)
    path = TEMPLATES / file_name(adr_type)
    sheet.save(path)
    return path


REGISTER = "Template_X-1.00_Registro_ADR_aperte.xlsx"
REGISTER_COLUMNS = ("ID dell'ADR", "Breve descrizione", "Team proponente", "Stato")


def make_register(rows=25):
    """The list of the open ADR: those not yet decided by the Design Authority (see storage.adr_register, which
    writes it filled in from the project)."""
    sheet = Sheet()
    ws = sheet.ws
    ws["A1"] = "X-1 · Registro delle decisioni architetturali (ADR) aperte"
    ws["A2"] = ws["A2"].value.replace("v1.0 · 28/09/2026", f"v1.1 · {date.today():%d/%m/%Y}")
    ws["A3"] = "Scopo. Tenere sotto controllo le ADR proposte e non ancora decise dalla Design Authority."
    ws["A4"] = ("Come si compila. Una riga per ADR aperta: l'ID e il titolo della sua scheda, il team che l'ha "
                "proposta e lo stato. Quando la Design Authority decide, l'ADR esce dal registro. In ATLAS lo stesso "
                "elenco, già compilato dal progetto, si ottiene con «Esporta ADR aperte…» nell'elenco X-1.")
    sheet.book.properties.title = "X-1 · Registro delle ADR aperte · template ARCADIA"
    register = Field("open_adrs", "ADR aperte", TABLE, columns=(
        Column("id", REGISTER_COLUMNS[0]), Column("title", REGISTER_COLUMNS[1]), Column("team", REGISTER_COLUMNS[2]),
        Column("status", REGISTER_COLUMNS[3], ENUM, options=STATUSES)))
    sheet.table(register, rows=rows, spans=("AA", "BD", "EE", "FF"))  # the description: the widest
    first = sheet.row - rows  # an example in the first row, in grey italics (to replace)
    for col, text in zip("ABEF", ("ADR-nnn", "Titolo della decisione", "es. T5 · Odometria", "Proposta")):
        cell = ws[f"{col}{first}"]
        cell.value = text
        cell.font = copy(sheet.hint_font)
    path = TEMPLATES / REGISTER
    sheet.save(path)
    return path


def main():
    # The register is written once, then kept as it is edited by hand («--register» writes it again).
    if "--register" in sys.argv or not (TEMPLATES / REGISTER).exists():
        print(make_register().name)
    for adr_type, when in ADR_TYPES:
        print(make(adr_type, when).name)


if __name__ == "__main__":
    main()
