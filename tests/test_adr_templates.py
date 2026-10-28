"""The Excel template of each type of ADR (Templates/, written by scripts/make_adr_templates.py): it has the fields
of the sheet of its type in ATLAS, so a change of the fields needs the templates written again."""
import importlib.util
from pathlib import Path

from openpyxl import load_workbook

from arcadia_tool.model.adr import common_fields, type_fields
from arcadia_tool.model.schema import ADR_TYPES

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("make_adr_templates", ROOT / "scripts" / "make_adr_templates.py")
make_adr_templates = importlib.util.module_from_spec(spec)
spec.loader.exec_module(make_adr_templates)


def test_each_type_has_its_template_with_the_fields_of_its_sheet():
    for adr_type, _ in ADR_TYPES:
        path = ROOT / "Templates" / make_adr_templates.file_name(adr_type)
        assert path.exists(), f"{path.name}: run scripts/make_adr_templates.py"
        ws = load_workbook(path).active
        labels = [ws[f"A{r}"].value for r in range(make_adr_templates.FIRST_ROW, ws.max_row + 1)]
        assert ws["A1"].value.endswith(adr_type)
        assert labels[0] == "Identificazione" and "Tipo di decisione" in labels
        expected = [f.label for f in type_fields(adr_type) + common_fields(adr_type)]
        assert [l for l in labels if l in expected] == sorted(expected, key=labels.index) and \
            set(expected) <= set(labels), f"{path.name}: run scripts/make_adr_templates.py"


def test_the_names_of_the_templates_follow_the_order_of_the_types():
    assert make_adr_templates.file_name("Nuovo dominio funzionale") == \
        "Template_X-1.01_ADR_Nuovo_dominio_funzionale.xlsx"
    assert make_adr_templates.file_name("Proprietà di un'informazione / confine tra domini") == \
        "Template_X-1.08_ADR_Proprieta_di_un_informazione_confine_tra_domini.xlsx"


def test_the_register_of_the_open_adr_has_its_columns():
    ws = load_workbook(ROOT / "Templates" / make_adr_templates.REGISTER).active
    # Edited by hand after it was written (see main): the columns are looked for in the whole sheet.
    header = next(r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == "ID dell'ADR")
    assert [ws[f"{c}{header}"].value for c in "ABEF"] == list(make_adr_templates.REGISTER_COLUMNS)
