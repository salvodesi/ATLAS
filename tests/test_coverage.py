"""The coverage report of the SRS items of board and of definition."""
import openpyxl
import pytest
from PySide6.QtWidgets import QApplication

from arcadia_tool.model.coverage import DERIVED, DOMAIN, FUNCTION, coverage, subset_coverage
from arcadia_tool.storage.project import Project
from arcadia_tool.ui.srs_reader import document_order


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def build(tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Acquisizione")
    d2 = p.new_element("domain", id="DF02", name="Odometria")
    n = iter(range(100))

    def item(ident, item_type, parent=None, category=None, **values):
        return p.new_element("requirement", id=ident, item_type=item_type, parent=parent.uuid if parent else None,
                             category=category, sequence=next(n), text=f"Testo {ident}", **values)

    s1 = item("3.4.1", "Sezione", title="Titolo 3.4.1", primary_domain=d1.uuid)
    by_section = item("3.4.1.1", "Requisito", s1, ["Bordo"])  # covered by the domain of its section
    own = item("3.4.1.2", "Requisito", s1, ["Bordo", "Terra"], contributing_domains=[d2.uuid])  # its own domains
    item("3.4.1.2.a", "Voce", own, ["Bordo"])  # a list item: counted with its clause, not alone
    last = item("3.4.1.2.1", "Requisito", own, ["Bordo"])  # a clause inside a clause: down to the last one
    s2 = item("3.5.1", "Sezione", title="Titolo 3.5.1")
    uncovered = item("3.5.1.1", "Requisito", s2, ["Bordo"])
    item("3.5.1.2", "Requisito", s2, ["Terra"])  # not in the report
    s3 = item("3.1.1", "Sezione", title="Definizioni")
    definition = item("3.1.1.1", "Requisito", s3, ["Definizione"], primary_domain=d2.uuid)
    undefined = item("3.1.1.2", "Requisito", s3, ["Definizione"])
    sf = p.new_element("system_function", id="SF-DF01-01", name="Acquisire", domain=d1.uuid, sections=[s1.uuid])
    derived = p.new_element("derived_requirement", id="DR-01", name="Derivato", parent=own.uuid)
    p.save()
    return p, dict(d1=d1, d2=d2, by_section=by_section, own=own, last=last, uncovered=uncovered,
                   definition=definition, undefined=undefined, sf=sf, derived=derived, s1=s1, s2=s2)


def test_an_item_is_covered_by_a_domain_and_also_by_functions_and_derived_requirements(tmp_path):
    p, e = build(tmp_path)
    ordered = document_order(p.model)
    board = coverage(p.model, ordered, "Bordo")
    rows = {r.item.get("id"): r for r in board.rows}
    assert list(rows) == ["3.4.1.1", "3.4.1.2", "3.4.1.2.1", "3.5.1.1"]  # the clauses, in the specification order
    assert rows["3.4.1.1"].domains == [(e["d1"], True)] and rows["3.4.1.1"].domains_from is e["s1"]
    assert rows["3.4.1.2"].domains == [(e["d2"], False)] and rows["3.4.1.2"].domains_from is None
    assert rows["3.4.1.2.1"].domains_from is e["own"] and rows["3.4.1.2.1"].parent is rows["3.4.1.2"]
    assert rows["3.4.1.1"].functions == [e["sf"]] and rows["3.4.1.2"].derived == [e["derived"]]
    assert not rows["3.5.1.1"].covered and board.uncovered == [rows["3.5.1.1"]]
    assert (board.count(DOMAIN), board.count(FUNCTION), board.count(DERIVED)) == (3, 3, 2)
    assert board.percent(DOMAIN) == pytest.approx(75.0)
    # Section by section.
    sections = board.sections()
    assert [(s.section.get("id"), s.total, s.count(DOMAIN)) for s in sections] == [("3.4.1", 3, 3), ("3.5.1", 1, 0)]
    assert rows["3.5.1.1"].section is e["s2"]
    definitions = coverage(p.model, ordered, "Definizione")
    assert [(r.item.get("id"), r.covered) for r in definitions.rows] == [("3.1.1.1", True), ("3.1.1.2", False)]
    assert subset_coverage([board, definitions]) == (4, 6, pytest.approx(200 / 3))


def test_the_report_opens_from_the_srs_page_with_sections_and_clauses_and_is_exported(app, tmp_path, monkeypatch):
    from arcadia_tool.ui import coverage_report as ui
    from arcadia_tool.ui.coverage_report import CATEGORY, STATE, CoverageReportDialog
    from arcadia_tool.ui.srs_reader import SrsReader

    p, e = build(tmp_path)
    reader = SrsReader(p)
    opened = []
    reader.open_element.connect(opened.append)
    assert reader.coverage_button.text() == "Report di copertura…"
    dialog = reader.show_coverage()
    assert isinstance(dialog, CoverageReportDialog) and not dialog.isModal()
    assert "66.7%" in dialog.subset.text() and "4 di 6" in dialog.subset.text()
    assert [dialog.tabs.tabText(i) for i in range(dialog.tabs.count())] == ["Bordo (3/4 coperte)",
                                                                           "Definizione (1/2 coperte)"]
    page = dialog.pages["Bordo"]
    # «Per sezione»: one row per section.
    assert [page.views.tabText(i) for i in range(page.views.count())] == ["Per sezione", "Per clausola"]
    table = page.by_section
    rows = [[table.topLevelItem(i).text(c) for c in range(7)] for i in range(table.topLevelItemCount())]
    assert rows == [["3.4.1", "Titolo 3.4.1", "3", "3 (100.0%)", "3 (100.0%)", "2 (66.7%)", "0"],
                    ["3.5.1", "Titolo 3.5.1", "1", "0 (0.0%)", "0 (0.0%)", "0 (0.0%)", "1"]]
    page.show_section(e["s2"].uuid)  # «Mostra le clausole della sezione» (right click on it)
    assert page.views.currentIndex() == 1 and page.table.currentItem().text(0) == "3.5.1"
    page._double_clicked(table.topLevelItem(0), 0)  # a double click on a section: its text, in a window
    assert dialog.viewer.isVisible() and dialog.viewer.element is e["s1"] and not opened
    # «Per clausola»: the sections, their clauses down to the last one, with their category.
    tree = page.table
    section = tree.topLevelItem(0)
    assert section.text(0) == "3.4.1" and section.text(STATE) == "3 di 3 coperte"
    assert [section.child(i).text(0) for i in range(section.childCount())] == ["3.4.1.1", "3.4.1.2"]
    clause = section.child(1)
    assert clause.text(CATEGORY) == "Bordo, Terra" and clause.child(0).text(0) == "3.4.1.2.1"
    assert tree.headerItem().text(CATEGORY) == "Categoria" and section.isExpanded() and clause.isExpanded()
    assert not tree.expandsOnDoubleClick()  # only the arrow opens and closes the sub-items
    page.filter.setCurrentIndex(4)  # «Senza requisito derivato»: 3.4.1.1 and 3.5.1.1
    visible = [tree.topLevelItem(i).isHidden() for i in range(tree.topLevelItemCount())]
    assert visible == [False, False] and clause.isHidden() and not section.child(0).isHidden()
    assert page.shown.text() == "2 di 4"
    page.filter.setCurrentIndex(0)
    page.search.setText("3.4.1.2.1")  # the clause inside a clause: its parent stays shown to reach it
    assert not clause.isHidden() and not clause.child(0).isHidden() and tree.topLevelItem(1).isHidden()
    page._double_clicked(clause.child(0), 0)  # a clause: its text in the window, not its sheet
    assert dialog.viewer.element is e["last"] and not opened
    assert dialog.viewer.sheet_button.isVisibleTo(dialog.viewer)
    dialog.viewer.sheet_button.click()  # from there, its sheet
    assert opened == [e["last"].uuid]
    # Exported: the summary, the sections and the clauses of each category.
    monkeypatch.setattr(ui.QMessageBox, "information", lambda *a, **k: None)
    path = dialog.export(tmp_path / "copertura.xlsx")
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Riepilogo", "Bordo - sezioni", "Bordo - clausole", "Definizione - sezioni",
                               "Definizione - clausole"]
    assert book["Riepilogo"]["B4"].value == "66.7% (4 di 6 voci)"
    sections = book["Bordo - sezioni"]
    assert [sections.cell(row=2, column=c).value for c in range(1, 8)] == [
        "3.4.1", "Titolo 3.4.1", 3, "3 (100.0%)", "3 (100.0%)", "2 (66.7%)", 0]
    clauses = book["Bordo - clausole"]
    assert [clauses.cell(row=r, column=1).value for r in range(2, 6)] == ["3.4.1.1", "3.4.1.2", "3.4.1.2.1",
                                                                         "3.5.1.1"]
    assert clauses.cell(row=3, column=3).value == "Bordo, Terra" and clauses.cell(row=3, column=2).value == "3.4.1"
    assert clauses.cell(row=5, column=8).value == "Scoperta" and "DR-01" in clauses.cell(row=3, column=7).value
    dialog.close()
