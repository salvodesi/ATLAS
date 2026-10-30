"""X-1 and X-2 as tables: a double click on the kind in the model tree opens a tab with the architectural decisions
or the requests for clarification, from which they are opened, created and deleted."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton

import arcadia_tool.ui.model_tree as model_tree
from arcadia_tool.storage.project import Project
from arcadia_tool.ui.element_list import ElementListPage
from arcadia_tool.ui.main_window import MainWindow
from arcadia_tool.ui.theme import apply_theme


@pytest.fixture(scope="module")
def app():
    app = QApplication.instance() or QApplication([])
    apply_theme(app)
    return app


@pytest.fixture
def recipients_asked(monkeypatch):
    """The window that asks the recipient of a new request: answered with the first one offered."""
    asked = []

    class Picker:
        def __init__(self, choices, selected, parent=None, single=False, title=""):
            asked.append((choices, set(selected)))
            self.choices = choices

        def exec(self):
            return True

        def selected(self):
            return [self.choices[0][0]]
    monkeypatch.setattr(model_tree, "RefPicker", Picker)
    return asked


def teams(p):
    t2 = p.new_element("team", name="Odometria")
    da = p.new_element("team", name="Architettura", design_authority=True)
    return da, t2


def buttons(page, row):
    return {b.text(): b for b in page.table.cellWidget(row, page.table.columnCount() - 1).findChildren(QPushButton)}


def test_double_click_on_x2_opens_the_table_of_the_requests(app, tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    da, t2 = teams(p)
    second = p.new_element("clarification", problem="Classificazione mancante", recipient=t2.uuid, status="Aperta")
    first = p.new_element("clarification", problem="Testo ambiguo", recipient=da.uuid, status="Chiusa")
    first.set("id", "CQ-01")
    second.set("id", "CQ-02")
    window = MainWindow()
    window.show()
    window._set_project(p)
    kind_item = window.tree._kind_items["clarification"]
    window.tree._double_clicked(kind_item, 0)
    page = window.list_pages["clarification"]
    assert window.pages.currentWidget() is page and not kind_item.isExpanded()  # a tab, not the branch
    table = page.table
    headers = [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())]
    assert headers == ["ID", "Destinatario della richiesta", "Stato", ""]
    rows = [[table.item(r, c).text() for c in range(3)] for r in range(table.rowCount())]
    assert rows == [["CQ-01", f"{da.get('id')} Architettura (Design Authority)", "Chiusa"],
                    ["CQ-02", f"{t2.get('id')} Odometria", "Aperta"]]
    assert "Testo ambiguo" in table.item(0, 0).toolTip()
    # The buttons are never cut: their column is as wide as they are.
    assert table.columnWidth(3) >= table.cellWidget(0, 3).sizeHint().width()

    buttons(page, 1)["Apri"].click()  # the sheet of the request in its tab
    assert window.pages.currentWidget() is window.editors[second.uuid]
    page._double_clicked(0, 1)  # a double click on the row too
    assert window.pages.currentWidget() is window.editors[first.uuid]

    first.set("status", "Aperta")  # changed in its sheet: the table follows
    p.update(first)
    window._element_changed(first)
    assert table.item(0, 2).text() == "Aperta"
    page.search.setText("odometria")
    assert table.isRowHidden(0) and not table.isRowHidden(1) and page.count.text() == "1 di 2"

    window.tree._double_clicked(kind_item, 0)  # opened again: the same tab
    assert list(window.list_pages.values()) == [page]
    window._set_project(None)  # closed with the project
    assert window.list_pages == {}
    window.close()


def test_requests_are_created_with_their_recipient_and_deleted(app, tmp_path, monkeypatch, recipients_asked):
    p = Project.create(tmp_path / "p", "Prova")
    window = MainWindow()
    window.show()
    window._set_project(p)
    window.open_element_list("clarification")
    page = window.list_pages["clarification"]
    # Empty: a message in the middle of the page with the button to create the first request, no search.
    assert page.body.currentWidget() is page.empty and page.search.isHidden()
    assert page.empty_text.text() == "Nessuna richiesta di chiarimento" and page.count.text() == "0 richieste"
    first = next(b for b in page.empty.findChildren(QPushButton))
    assert first.text() == "Nuova richiesta di chiarimento"
    first.click()  # no team yet: a recipient is needed first
    assert not p.model.of_kind("clarification") and not recipients_asked

    da, t2 = teams(p)
    first.click()  # the recipient is asked: the Design Authority first (and proposed), then the teams
    [(choices, proposed)] = recipients_asked
    assert [label for _, label in choices] == [p.model.label(da), p.model.label(t2)] and proposed == {da.uuid}
    [request] = p.model.of_kind("clarification")
    assert request.get("id") == "CQ-01" and request.get("recipient") == da.uuid
    assert window.editors[request.uuid].editing  # its sheet opened in edit mode
    assert page.body.currentWidget() is page.table and page.table.item(0, 1).text().endswith("(Design Authority)")
    page.new_button.click()  # «Nuova richiesta», always in the header
    assert len(p.model.of_kind("clarification")) == 2 and page.count.text() == "2 richieste"
    window.pages.setCurrentWidget(page)  # back to the table

    assert set(buttons(page, 0)) == {"Apri", "Elimina"}  # two buttons of the same kind (same style)
    buttons(page, 0)["Elimina"].click()  # asked to confirm: «No» (tests)
    assert len(p.model.of_kind("clarification")) == 2
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    buttons(page, 0)["Elimina"].click()
    assert [e.get("id") for e in p.model.of_kind("clarification")] == ["CQ-02"]
    assert page.table.rowCount() == 1 and window.pages.currentWidget() is page
    assert request.uuid not in window.editors  # its open sheet closed with it
    window.close()


def test_the_clause_of_a_request_is_chosen_with_ticks(app, tmp_path, monkeypatch):
    import arcadia_tool.ui.srs_trace as srs_trace
    from arcadia_tool.ui.element_editor import ElementEditor
    p = Project.create(tmp_path / "p", "Prova")
    section = p.new_element("requirement", id="3.15.1", item_type="Sezione")
    clause = p.new_element("requirement", id="3.15.1.3.9", item_type="Requisito", parent=section.uuid)
    request = p.new_element("clarification", problem="Nessuna categoria")
    assert request.schema.field("clause").picker  # like the other references to the specification
    used = []

    class Ticks(srs_trace.TraceClauseDialog):
        def exec(self):
            used.append(type(self))
            self.set_state(clause.uuid, srs_trace.PARENT)  # the clause ticked
            self.changed = self.apply()
            return True
    monkeypatch.setitem(srs_trace.ITEM_DIALOGS, "clause", Ticks)
    editor = ElementEditor()
    editor.edit(p, request, editing=True)
    editor._pick(request.schema.field("clause"))
    assert used == [Ticks] and editor.draft.get("clause") == clause.uuid
    editor.save_changes()
    assert request.get("clause") == clause.uuid


def test_recipients_written_as_text_become_teams(tmp_path):
    from arcadia_tool.model.serialization import dumps
    p = Project.create(tmp_path / "p", "Prova")
    da, t2 = teams(p)
    texts = ["DA", t2.get("id"), "odometria", "ERA / UNISIG"]
    requests = [p.new_element("clarification", problem=f"P{i}") for i in range(4)]
    for request, text in zip(requests, texts):
        request.values["recipient"] = text  # as earlier versions saved it
    p.save()
    reopened = Project.open(tmp_path / "p")
    recipients = [reopened.model.get(r.uuid).get("recipient") for r in requests]
    assert recipients == [da.uuid, t2.uuid, t2.uuid, None]
    assert reopened.model.get(requests[3].uuid).get("recipient_text") == "ERA / UNISIG"  # kept, shown in the table
    page = ElementListPage(reopened, "clarification")
    row = next(r for r in range(page.table.rowCount()) if page.table.item(r, 1).text() == "ERA / UNISIG")
    assert row is not None


def test_x1_lists_the_architectural_decisions(app, tmp_path, monkeypatch):
    p = Project.create(tmp_path / "p", "Prova")
    team = p.new_element("team", id="T2", name="Odometria", members=[{"name": "Luca Neri"}])
    adr = p.new_element("adr", title="MRSP in un dominio autonomo", status="Decisa",
                        context="La catena rimbalza tra DF03 e DF04.", proposer=team.uuid,
                        proposer_member="Luca Neri")
    adr.set("decision_date", "12/10/2026")  # not in the table
    p.update(adr)
    window = MainWindow()
    window.show()
    window._set_project(p)
    window.tree._double_clicked(window.tree._kind_items["adr"], 0)
    page = window.list_pages["adr"]
    table = page.table
    assert [table.horizontalHeaderItem(c).text() for c in range(7)] == [
        "ID", "Titolo", "Tipo", "DF impattati", "Team proponente", "Membro proponente", "Stato"]
    assert [table.item(0, c).text() for c in range(7)] == ["ADR-001", "MRSP in un dominio autonomo", "", "",
                                                           "T2 Odometria", "Luca Neri", "Decisa"]
    assert "rimbalza" in table.item(0, 0).toolTip() and page.count.text() == "1 ADR"
    import arcadia_tool.ui.adr_types as adr_types

    def choose(dialog):  # the type of the new ADR, chosen in its window
        dialog.adr_type = "Fusione di domini"
        return True
    monkeypatch.setattr(adr_types.AdrTypeDialog, "exec", choose)
    page.new_button.click()  # the type asked, then created and opened
    new = next(e for e in p.model.of_kind("adr") if e.uuid != adr.uuid)
    assert new.get("id") == "ADR-002" and window.editors[new.uuid].editing
    assert new.get("adr_type") == "Fusione di domini" and new.get("status") == "Proposta"
    window.pages.setCurrentWidget(page)
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    rows = {page.table.item(r, 0).text(): r for r in range(page.table.rowCount())}
    assert list(rows) == ["ADR-001", "ADR-002"]  # by ID
    assert page.table.item(rows["ADR-002"], 2).text() == "Fusione di domini"
    buttons(page, rows["ADR-002"])["Elimina"].click()
    assert [e.get("id") for e in p.model.of_kind("adr")] == ["ADR-001"]
    window.close()


def test_the_list_of_the_adr_is_sorted_by_id_and_by_a_click_on_a_heading(app, tmp_path):
    from PySide6.QtCore import Qt

    p = Project.create(tmp_path / "p", "Prova")
    t1 = p.new_element("team", id="T1", name="Odometria", members=[{"name": "Luca Neri"}, {"name": "Anna Bianchi"}])
    t2 = p.new_element("team", id="T2", name="Fondazione", members=[{"name": "Marta Verdi"}])
    for title, team, member, adr_type, status, decision in (
            ("Prima", t2, "Marta Verdi", "Eliminazione di dominio", "Decisa", "Respinta"),
            ("Seconda", t1, "Luca Neri", "Altro / decisione trasversale", "Proposta", None),
            ("Terza", None, None, None, "Superata", None),
            ("Quarta", t1, "Anna Bianchi", "Riallocazione di sezione o clausola", "Revisionato DA", "Accolta")):
        p.new_element("adr", title=title, status=status, **({"proposer": team.uuid} if team else {}),
                      **({"proposer_member": member} if member else {}),
                      **({"adr_type": adr_type} if adr_type else {}), **({"da_decision": decision} if decision else {}))
    for n in range(5, 11):  # ADR-010 after ADR-009, not after ADR-001
        p.new_element("adr", title=f"Altra {n}", status="Proposta")
    page = ElementListPage(p, "adr")

    def ids():
        return [page.table.item(r, 0).text() for r in range(page.table.rowCount())]

    def column(name):
        c = page.columns.index(name)
        return [page.table.item(r, c).text() for r in range(page.table.rowCount())]
    assert ids() == [f"ADR-{n:03d}" for n in range(1, 11)]
    # The decision of the DA next to the status of the ADR, in colour like it.
    status, decision = page.columns.index("status"), page.columns.index("da_decision")
    assert decision == status + 1 and page.table.horizontalHeaderItem(decision).text() == "Decisione DA"
    assert column("da_decision")[:4] == ["Respinta", "", "", "Accolta"]
    assert page.table.item(0, decision).foreground().color().name().upper() == "#B42318"
    assert page.table.item(3, decision).foreground().color().name().upper() == "#166534"
    assert column("status")[3] == "Revisionato DA"
    page.sort_by("da_decision")  # in the order of the outcomes: Accolta, …, Respinta; the others last
    assert column("da_decision")[:2] == ["Accolta", "Respinta"] and ids()[:2] == ["ADR-004", "ADR-001"]
    page.sort_by("proposer")  # the teams, the ADR without one last (by ID)
    assert column("proposer")[:4] == ["T1 Odometria", "T1 Odometria", "T2 Fondazione", ""]
    assert ids()[:4] == ["ADR-002", "ADR-004", "ADR-001", "ADR-003"]
    page.sort_by("proposer_member")
    assert column("proposer_member")[:3] == ["Anna Bianchi", "Luca Neri", "Marta Verdi"]
    page.sort_by("adr_type", Qt.DescendingOrder)  # in the order of the types of the method, here reversed
    assert column("adr_type")[:3] == ["Altro / decisione trasversale", "Riallocazione di sezione o clausola",
                                      "Eliminazione di dominio"]
    page.sort_by("status")  # in the order of the states
    assert column("status")[0] == "Proposta" and column("status")[-3:] == ["Revisionato DA", "Decisa", "Superata"]
    page.sort_by("title")  # not sortable: the order stays, as the arrow
    assert column("status")[-3:] == ["Revisionato DA", "Decisa", "Superata"]
    header = page.table.horizontalHeader()
    assert header.sortIndicatorSection() == page.columns.index("status")
    from PySide6.QtTest import QTest
    page.resize(1200, 600)
    page.show()
    id_column = page.columns.index("id")
    point = header.rect().center()
    point.setX(header.sectionViewportPosition(id_column) + header.sectionSize(id_column) // 2)
    QTest.mouseClick(header.viewport(), Qt.LeftButton, Qt.NoModifier, point)  # a real click: by ID
    assert ids()[0] in ("ADR-001", "ADR-010") and header.sortIndicatorSection() == id_column
    if ids()[0] == "ADR-001":
        QTest.mouseClick(header.viewport(), Qt.LeftButton, Qt.NoModifier, point)  # again: the other way round
    assert ids()[0] == "ADR-010"
    page.sort_by("id", Qt.DescendingOrder)
    assert ids()[0] == "ADR-010"
    p.new_element("adr", title="Nuova", status="Proposta")
    page.reload()  # a new ADR: the order chosen is kept
    assert ids()[0] == "ADR-011"


def test_the_system_functions_open_as_a_table_with_their_status_on_the_row(app, tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Protezione")
    clause = p.new_element("requirement", id="3.4.1.1", item_type="Requisito")
    sf = p.new_element("system_function", name="Proteggere dal Roll Away", domain=d1.uuid, sections=[clause.uuid],
                       status="Proposta")
    p.new_element("derived_requirement", parent=clause.uuid, obligation="Frena.", system_function=sf.uuid)
    p.save()
    window = MainWindow()
    window.show()
    window._set_project(p)
    window.tree._double_clicked(window.tree._kind_items["system_function"], 0)  # a tab, not the branch
    page = window.list_pages["system_function"]
    table = page.table
    assert [table.horizontalHeaderItem(c).text() for c in range(6)] == [
        "ID", "Nome (verbo + oggetto)", "Dominio", "Sezioni e clausole SRS", "Requisiti derivati", "Stato"]
    assert [table.item(0, c).text() for c in range(5)] == ["SF-DF01-01", "Proteggere dal Roll Away",
                                                           "DF01 Protezione", "1", "1"]
    choice = table.cellWidget(0, 5)  # the status chosen on the row, saved at once
    assert "Elimina" not in [b.text() for b in table.cellWidget(0, 6).findChildren(QPushButton)]
    choice.setCurrentText("Rivista")
    app.processEvents()
    assert sf.get("status") == "Rivista" and sf.uuid not in p.pending()
    window.close()


def test_the_list_stays_where_it_is_when_a_status_is_chosen(app, tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Protezione")
    functions = [p.new_element("system_function", name=f"Funzione {n:02d}", domain=d1.uuid, status="Proposta")
                 for n in range(1, 41)]
    p.save()
    window = MainWindow()
    window.resize(1200, 600)
    window.show()
    window._set_project(p)
    window.tree._double_clicked(window.tree._kind_items["system_function"], 0)
    page = window.list_pages["system_function"]
    table, status = page.table, page.columns.index("status")
    app.processEvents()
    row = 30
    table.scrollToItem(table.item(row, 0))
    app.processEvents()
    scroll = table.verticalScrollBar().value()
    assert scroll > 0
    table.cellWidget(row, status).setCurrentText("Rivista")  # saved, the list rebuilt
    app.processEvents()
    app.processEvents()
    uuid = table.item(row, 0).data(Qt.UserRole)
    assert p.model.get(uuid).get("status") == "Rivista"
    assert table.currentRow() == row and table.verticalScrollBar().value() == scroll  # still there
    assert table.cellWidget(row, status).currentText() == "Rivista"
    window._set_project(None)
    window.close()


def test_the_list_of_the_adr_shows_the_domains_they_are_about(app, tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Protezione")
    d2 = p.new_element("domain", id="DF02", name="Localizzazione")
    section = p.new_element("requirement", id="3.4.1", item_type="Sezione", primary_domain=d1.uuid)
    clause = p.new_element("requirement", id="3.4.1.1", item_type="Requisito", parent=section.uuid)
    p.new_element("adr", title="Spostare", adr_type="Riallocazione di sezione o clausola", status="Proposta",
                  realloc_item=[clause.uuid], primary_proposed=d2.uuid)  # DF01: from its section, DF02: proposed
    p.new_element("adr", title="Eliminare", adr_type="Eliminazione di dominio", status="Proposta",
                  deleted_domain=d2.uuid)
    p.new_element("adr", title="Altro", status="Proposta")
    page = ElementListPage(p, "adr")
    c = page.columns.index("adr_domains")
    assert page.table.horizontalHeaderItem(c).text() == "DF impattati"
    assert [page.table.item(r, c).text() for r in range(3)] == ["DF01, DF02", "DF02", ""]
    page.sort_by("adr_domains")
    assert [page.table.item(r, c).text() for r in range(3)] == ["DF01, DF02", "DF02", ""]


@pytest.mark.parametrize("kind, text", [("clarification", "0 richieste"), ("adr", "0 ADR")])
def test_empty_list(app, tmp_path, kind, text):
    p = Project.create(tmp_path / "p", "Prova")
    page = ElementListPage(p, kind)
    assert page.table.rowCount() == 0 and page.body.currentWidget() is page.empty
    assert page.count.text() == text


def test_a_new_element_never_saved_goes_away(app, tmp_path, monkeypatch):
    import arcadia_tool.ui.adr_types as adr_types
    from arcadia_tool.ui.model_tree import UUID_ROLE

    p = Project.create(tmp_path / "p", "Prova")
    kept = p.new_element("domain", name="Protezione")
    p.save()
    window = MainWindow()
    window._set_project(p)

    def tree_uuids():
        found, stack = set(), [window.tree.tree.topLevelItem(i) for i in range(window.tree.tree.topLevelItemCount())]
        while stack:
            item = stack.pop()
            if item.data(0, UUID_ROLE):
                found.add(item.data(0, UUID_ROLE))
            stack.extend(item.child(i) for i in range(item.childCount()))
        return found

    # «Nuovo» then «Annulla»: the domain is not there any more, nothing to save
    window.tree.create("domain")
    new = next(d for d in p.model.of_kind("domain") if d is not kept)
    editor = window.editors[new.uuid]
    assert editor.editing and editor.unsaved_new and new.uuid in tree_uuids()
    editor.cancel_changes(ask=False)
    assert p.model.get(new.uuid) is None and new.uuid not in tree_uuids() and new.uuid not in window.editors
    assert not p.is_modified

    # an ADR: its tab closed without saving
    monkeypatch.setattr(adr_types.AdrTypeDialog, "exec", lambda dialog: dialog._accept() or True)
    window.tree.create("adr")
    adr = p.model.of_kind("adr")[0]
    window._close_page(window.pages.indexOf(window.editors[adr.uuid]))
    assert p.model.of_kind("adr") == [] and adr.uuid not in tree_uuids()

    # saved: it stays, also after «Modifica» and «Annulla»
    window.tree.create("domain")
    saved = next(d for d in p.model.of_kind("domain") if d is not kept)
    editor = window.editors[saved.uuid]
    editor.save_changes()
    editor.start_editing()
    editor.cancel_changes(ask=False)
    assert p.model.get(saved.uuid) is saved and saved.uuid in tree_uuids()
    window._set_project(None)
    window.close()


def test_no_recovery_is_left_when_the_edits_were_undone(app, tmp_path):
    from arcadia_tool.storage.project import recovery_file

    p = Project.create(tmp_path / "p", "Prova")
    window = MainWindow()
    window._set_project(p)
    window.tree.create("domain")  # a new element: the project is modified, its copy for the recovery written
    new = p.model.of_kind("domain")[0]
    p.write_recovery()
    assert recovery_file(p.path).exists()
    window.editors[new.uuid].cancel_changes(ask=False)  # cancelled: nothing changed any more
    assert not p.is_modified and not recovery_file(p.path).exists()

    p.new_element("team", name="Radio")  # modified, copy written, then the edit undone in another way
    p.write_recovery()
    p.discard_new(p.model.of_kind("team")[0].uuid)
    assert window._confirm_discard()  # a normal close: no recovery at the next start
    assert not recovery_file(p.path).exists()
    window._set_project(None)
    window.close()


def test_the_open_adr_are_exported_to_excel(app, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from arcadia_tool.storage.adr_register import open_adrs

    p = Project.create(tmp_path / "p", "EVC")
    team = p.new_element("team", id="T5", name="Odometria")
    p.new_element("adr", title="Stimare la velocità in DF07", status="Proposta", proposer=team.uuid)
    p.new_element("adr", title="Senza stato")
    p.new_element("adr", title="Già decisa", status="Decisa")
    p.new_element("adr", title="Superata", status="Superata")
    assert [(r[1], r[2], r[3]) for r in open_adrs(p.model)] == [
        ("Stimare la velocità in DF07", p.model.label(team), "Proposta"), ("Senza stato", "", "Proposta")]
    page = ElementListPage(p, "adr")
    out = tmp_path / "aperte.xlsx"
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    assert page.export_open(str(out)) == 2
    ws = load_workbook(out).active
    assert [ws.cell(4, c).value for c in range(1, 5)] == ["ID", "Descrizione", "Team proponente", "Stato"]
    assert [ws.cell(5, c).value for c in range(2, 5)] == ["Stimare la velocità in DF07", p.model.label(team),
                                                          "Proposta"]
    assert ws.cell(5, 1).value.startswith("ADR-") and ws.max_row == 6
    p.close()


def test_ctrl_s_saves_the_changes_of_the_sheet_shown(app, tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    window = MainWindow()
    window._set_project(p)
    window.tree.create("domain")
    domain = p.model.of_kind("domain")[0]
    editor = window.editors[domain.uuid]
    editor.save_changes()
    assert not p.is_modified and not window.save_action.isEnabled()
    editor.start_editing()
    editor.widgets["name"].setText("Odometria")
    editor.flush()
    assert editor.dirty and window.save_action.isEnabled()  # only the sheet has changes: Ctrl+S works
    window.save_action.trigger()
    assert not editor.dirty and p.model.get(domain.uuid).get("name") == "Odometria" and not p.is_modified
    window._set_project(None)
    window.close()


def test_double_click_on_sa1_opens_the_table_of_the_domains_with_their_status_on_the_rows(app, tmp_path):
    from PySide6.QtWidgets import QComboBox

    p = Project.create(tmp_path / "p", "Prova")
    da, t2 = teams(p)
    t2.set("referent", "Mario Rossi")
    second = p.new_element("domain", id="DF02", name="Odometria", team=t2.uuid, status="Bozza")
    first = p.new_element("domain", id="DF01", name="Supervisione", mission="Proteggere il treno")
    p.save()
    window = MainWindow()
    window.show()
    window._set_project(p)
    kind_item = window.tree._kind_items["domain"]
    window.tree._double_clicked(kind_item, 0)
    page = window.list_pages["domain"]
    assert window.pages.currentWidget() is page and not kind_item.isExpanded()  # a tab, not the branch
    table = page.table
    headers = [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())]
    assert headers == ["ID", "Nome del dominio", "Team assegnatario", "Responsabile", "Stato", ""]
    rows = [[table.item(r, c).text() for c in range(4)] for r in range(table.rowCount())]
    assert rows == [["DF01", "Supervisione", "", ""], ["DF02", "Odometria", f"{t2.get('id')} Odometria", "Mario Rossi"]]
    assert "Proteggere il treno" in table.item(0, 0).toolTip()
    assert list(buttons(page, 0)) == ["Apri"]  # only the button to its sheet
    buttons(page, 1)["Apri"].click()
    assert window.pages.currentWidget() is window.editors[second.uuid]
    # The status chosen on the row: in its colour, in the model and on disk at once (only it).
    choice = table.cellWidget(1, 4)
    assert isinstance(choice, QComboBox) and choice.currentText() == "Bozza"
    from arcadia_tool.ui.theme import status_colors
    assert status_colors("Bozza")[0] in choice.styleSheet()
    other = p.new_element("team", name="Radio")  # a change made elsewhere, not saved
    choice.setCurrentText("Approvato")
    app.processEvents()
    assert second.get("status") == "Approvato" and second.uuid not in p.changes_since(0)
    assert other.uuid in p.changes_since(0)
    choice = table.cellWidget(1, 4)  # the table drawn again
    assert choice.currentText() == "Approvato" and status_colors("Approvato")[0] in choice.styleSheet()
    assert window.editors[second.uuid].element.get("status") == "Approvato"
    window._set_project(None)
    window.close()
