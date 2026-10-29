"""«Importa clausole principali…» of a domain sheet: the row «Sezioni SRS principali» of an Excel SA-1 sheet."""
import openpyxl
import pytest
from PySide6.QtWidgets import QApplication, QPushButton

from arcadia_tool.model.srs_sections import expand, parse_sections, resolve
from arcadia_tool.storage.domain_sheets import read_domain_sheets, sheet_of
from arcadia_tool.storage.project import Project


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def test_single_sections_groups_and_other_documents():
    refs = resolve(parse_sections("3.4.1, A.3.3, 4.8.1–4.8.5, 3.6.3.2; SUBSET-034/119\nERA_ERTMS_015560, "
                                  "A.3.7-A.3.10"))
    assert [(r.text, r.ids, r.external) for r in refs] == [
        ("3.4.1", ["3.4.1"], False), ("A.3.3", ["A.3.3"], False),
        ("4.8.1–4.8.5", ["4.8.1", "4.8.2", "4.8.3", "4.8.4", "4.8.5"], False),
        ("3.6.3.2", ["3.6.3.2"], False), ("SUBSET-034/119", [], True), ("ERA_ERTMS_015560", [], True),
        ("A.3.7-A.3.10", ["A.3.7", "A.3.8", "A.3.9", "A.3.10"], False)]
    # A group across two parents: the sections between them in the order of the specification, at the same depth.
    order = ["3.13.10", "3.13.11", "3.14", "3.14.1", "3.14.1.1", "3.14.2"]
    assert expand("3.13.10", "3.14.2", order) == ["3.13.10", "3.13.11", "3.14.1", "3.14.2"]


def workbook(path):
    wb = openpyxl.Workbook()
    index = wb.active
    index.title = "Indice"
    index.append(["ID", "Domain"])
    for ident, name, principal in (("DF01", "Acquisizione", "3.4.1, 4.8.1–4.8.5, SUBSET-027"),
                                   ("DF02", "Odometria", "3.6.1, 4.8.3")):  # 4.8.3: also in DF01
        ws = wb.create_sheet(ident)
        ws.append([f"{ident} {name}"])
        ws.append(["Template", "SA-1"])
        ws.append(["ID e nome", f"{ident} {name}"])
        ws.append(["Sezioni SRS principali", principal])
    wb.save(path)
    return path


def test_the_sheets_of_the_domains_are_read_from_the_workbook(tmp_path):
    sheets = read_domain_sheets(workbook(tmp_path / "sa1.xlsx"))
    assert [(s.sheet, s.ident, s.name, s.principal) for s in sheets] == [
        ("DF01", "DF01", "Acquisizione", "3.4.1, 4.8.1–4.8.5, SUBSET-027"), ("DF02", "DF02", "Odometria", "3.6.1, 4.8.3")]
    p = Project.create(tmp_path / "p", "Prova")
    assert sheet_of(sheets, p.new_element("domain", id="DF02", name="Altro nome")).sheet == "DF02"
    assert sheet_of(sheets, p.new_element("domain", name="Acquisizione")).sheet == "DF01"  # by its name
    assert sheet_of(sheets, p.new_element("domain", id="DF09", name="Nessuno")) is None


def test_a_sheet_with_the_id_and_name_of_another_domain_is_read_by_its_name_and_noted(tmp_path):
    wb = openpyxl.Workbook()
    for ident, title, heading, principal in (("DF13", "DF13 National Systems Interaction",
                                              "DF14 Stored Data Management", "3.15.6, SUBSET-035"),
                                             ("DF14", "DF14 Stored Data Management", "DF14 Stored Data Management",
                                              "3.18.1")):
        ws = wb.create_sheet(ident)
        ws.append([title])
        ws.append(["ID e nome", heading])
        ws.append(["Sezioni SRS principali", principal])
    wb.save(tmp_path / "sa1.xlsx")
    first, second = read_domain_sheets(tmp_path / "sa1.xlsx")
    assert (first.ident, first.name, second.ident) == ("DF13", "National Systems Interaction", "DF14")
    assert "DF14 Stored Data Management" in first.note and not second.note
    p = Project.create(tmp_path / "p", "Prova")
    assert sheet_of([first, second], p.new_element("domain", id="DF14", name="Stored Data Management")) is second


def project_with_items(tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Acquisizione")
    d2 = p.new_element("domain", id="DF02", name="Odometria")
    items = {}
    for i, ident in enumerate(("3.4.1", "4.8.1", "4.8.2", "4.8.3", "4.8.5", "3.6.1")):
        items[ident] = p.new_element("requirement", id=ident, title=f"Titolo {ident}", item_type="Sezione",
                                     category=["Bordo"], sequence=i)
    items["4.8.2"].set("primary_domain", d2.uuid)  # another domain is principal
    items["4.8.1"].set("primary_domain", d1.uuid)  # already principal
    items["3.6.1"].set("contributing_domains", [d1.uuid])  # a contribution: it stays
    p.save()
    return p, d1, d2, items


def test_the_summary_shows_the_sections_mapped_and_the_import_makes_the_domain_principal(app, tmp_path):
    from PySide6.QtCore import Qt
    from arcadia_tool.ui.import_principal import ALREADY, EXTERNAL, MISSING, NEW, REPLACED, ImportPrincipalDialog

    p, d1, d2, items = project_with_items(tmp_path)
    dialog = ImportPrincipalDialog(p, d1, workbook(tmp_path / "sa1.xlsx"))
    assert dialog.sheet.currentData().sheet == "DF01" and dialog.warning.isHidden()  # its sheet, by its ID
    assert [(r[0], r[2], r[3]) for r in dialog.rows] == [
        (NEW, "3.4.1", "3.4.1"), (ALREADY, "4.8.1", "4.8.1–4.8.5"), (REPLACED, "4.8.2", "4.8.1–4.8.5"),
        (NEW, "4.8.3", "4.8.1–4.8.5"), (MISSING, "4.8.4", "4.8.1–4.8.5"), (NEW, "4.8.5", "4.8.1–4.8.5"),
        (EXTERNAL, "SUBSET-027", "SUBSET-027")]
    first = dialog.table.topLevelItem(0)
    assert [first.text(c) for c in range(1, 5)] == ["3.4.1", "Titolo 3.4.1", "Sezione", "Bordo"]
    from arcadia_tool.ui.import_principal import ALSO, TICK
    assert dialog.table.headerItem().text(TICK) == "Importa"
    # Every item to import has its tick, on at first; the others none.
    entries = [dialog.table.topLevelItem(i) for i in range(len(dialog.rows))]
    assert [e.data(TICK, Qt.CheckStateRole) is not None for e in entries] == [r.importable for r in dialog.rows]
    assert all(e.checkState(TICK) == Qt.Checked for e, r in zip(entries, dialog.rows) if r.importable)
    # Allocated to more than one domain: principal of DF02 in the model (4.8.2), in the sheet of DF02 too (4.8.3).
    assert [r.ident for r in dialog.rows if r.overlapping] == ["4.8.2", "4.8.3"]
    replaced, listed = dialog.table.topLevelItem(2), dialog.table.topLevelItem(3)
    assert "DF02" in replaced.text(ALSO) and "principale nel modello" in replaced.text(ALSO)
    assert listed.text(ALSO) == "DF02 (Excel)"
    assert dialog.counts.text().startswith("4 da importare") and "2 con la spunta" in dialog.overlaps.text()
    listed.setCheckState(TICK, Qt.Unchecked)  # not imported
    first.setCheckState(TICK, Qt.Unchecked)  # also one allocated only to this domain
    assert dialog.counts.text().startswith("2 da importare") and "1 con la spunta" in dialog.overlaps.text()
    dialog.import_button.click()
    assert all(items[i].get("primary_domain") == d1.uuid for i in ("4.8.1", "4.8.2", "4.8.5"))
    assert items["4.8.3"].get("primary_domain") is None and items["3.4.1"].get("primary_domain") is None
    assert items["3.6.1"].get("contributing_domains") == [d1.uuid]
    assert {e.uuid for e in dialog.changed} == {items[i].uuid for i in ("4.8.2", "4.8.5")}
    # Another sheet chosen by hand.
    other = ImportPrincipalDialog(p, p.new_element("domain", id="DF09", name="Nuovo"), tmp_path / "sa1.xlsx")
    assert not other.warning.isHidden()
    other.sheet.setCurrentIndex(1)
    assert [(r.outcome, r.ident, r.elsewhere) for r in other.rows] == [(NEW, "3.6.1", ()), (NEW, "4.8.3", ("DF01",))]


def test_the_domain_sheet_imports_and_its_save_writes_the_sections(app, tmp_path, monkeypatch):
    from arcadia_tool.ui import import_principal as ui
    from arcadia_tool.ui.main_window import MainWindow

    p, d1, d2, items = project_with_items(tmp_path)
    path = workbook(tmp_path / "sa1.xlsx")
    window = MainWindow()
    window._set_project(p)
    window._select(d1)
    editor = window.editor
    texts = [b.text() for b in editor.widget().findChildren(QPushButton)]
    assert "Traccia sezioni e clausole SRS…" in texts and "Importa clausole principali…" in texts
    monkeypatch.setattr(ui.ImportPrincipalDialog, "exec", lambda dialog: dialog.import_sections())
    elsewhere = p.new_element("team", name="Radio")  # not saved by the sheet
    changed = editor._import_principal(path)
    assert len(changed) == 4 and items["4.8.3"].get("primary_domain") == d1.uuid
    window.save_tab()  # Ctrl+S on the sheet: the sections imported are written, only them
    assert not {items[i].uuid for i in ("3.4.1", "4.8.2", "4.8.3", "4.8.5")} & p.pending()
    assert elsewhere.uuid in p.pending()
    window._set_project(None)
    window.close()
