"""«Genera documentazione»: the artefacts SA-1 … SA-7 and the ADR in Excel, filled in from the model with the layout
of their templates (Templates/, written by scripts/make_sa_templates.py from the same code)."""
import importlib.util
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

from arcadia_tool.documents.generate import generate, make_doc, subjects
from arcadia_tool.documents.sa import ARTIFACTS, Doc, write_sheet
from arcadia_tool.storage.domain_sheets import read_domain_sheets
from arcadia_tool.storage.function_sheets import read_function_sheets
from arcadia_tool.storage.project import Project

ROOT = Path(__file__).resolve().parent.parent
TODAY = date(2026, 10, 30)


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def values(ws):
    """The texts of the sheet, row by row (the empty cells left out)."""
    return [[str(c) for c in row if c not in (None, "")] for row in ws.iter_rows(values_only=True)]


def find(ws, label):
    """The value next to a label in column A."""
    for row in ws.iter_rows(values_only=True):
        if row[0] == label:
            return row[1]
    raise KeyError(label)


def table(ws, heading):
    """The rows under the heading row whose first cell is `heading`, up to the first empty row."""
    rows, inside = [], False
    for row in ws.iter_rows(values_only=True):
        if not inside:
            inside = row[0] == heading
            header = [c for c in row]
            continue
        if not any(row):
            break
        rows.append({h: v for h, v in zip(header, row) if h})
    return rows


def build(tmp_path):
    p = Project.create(tmp_path / "p", "Prova SA")
    t1 = p.new_element("team", id="T1", name="Odometria", members=[{"name": "Anna Bianchi"}], referent="Anna Bianchi")
    d1 = p.new_element("domain", id="DF01", name="Odometria", mission="Stimare la **velocità**.", team=t1.uuid,
                       status="In revisione")
    d2 = p.new_element("domain", id="DF02", name="Frenatura", status="Bozza")
    s1 = p.new_element("requirement", id="3.4.1", item_type="Sezione", sequence=1, primary_domain=d1.uuid)
    c1 = p.new_element("requirement", id="3.4.1.1", item_type="Requisito", parent=s1.uuid, sequence=2,
                       category=["Bordo"], contributing_domains=[d2.uuid])
    c2 = p.new_element("requirement", id="3.4.1.2", item_type="Requisito", parent=s1.uuid, sequence=3,
                       category=["Bordo"])
    sf1 = p.new_element("system_function", name="Stimare la velocità", domain=d1.uuid, sections=[s1.uuid],
                        rows_452=["4.5.2-R01"], mode_availability={"FS": "X", "SR": "O"}, is_new=True,
                        new_reason="Manca nel catalogo", status="Proposta")
    sf2 = p.new_element("system_function", name="Comandare i freni", domain=d2.uuid, sections=[c1.uuid])
    lf = p.new_element("logical_function", name="Calcolare la frenata", refines=sf2.uuid)  # of DF02
    ei = p.new_element("exchange_item", id="EI-velocita", name="Velocità stimata", type="Dato logico", nature="FLOW",
                       elements=[{"variable": "V_TRAIN", "unit": "km/h"}])
    fe = p.new_element("functional_exchange", name="Velocità", source=sf1.uuid, target=sf2.uuid,
                       exchange_item=ei.uuid, nature="FLOW", production_condition="Sempre", clauses=[c1.uuid])
    derived = p.new_element("derived_requirement", name="Frenare", parent=c1.uuid, obligation="Il treno frena.",
                            function=lf.uuid, domain=d2.uuid)
    p.new_element("trace_record", requirement=c1.uuid, primary_function=sf1.uuid, secondary_functions=[sf2.uuid],
                  exchange_items=[derived.uuid], outcome="Tracciata a funzione", status="Proposto")
    chain = p.new_element("functional_chain", id="FC-frenatura", title="Frenatura", responsible_domain=d2.uuid,
                          steps=[{"cycle": "C0", "function": sf1.uuid, "what": "Stima", "data": "Velocità",
                                  "nature": "FLOW", "exchange": fe.uuid, "srs": "3.4.1.1"},
                                 {"cycle": "C0", "function": sf2.uuid, "what": "Frena"}])
    realloc = p.new_element("adr", title="Spostare 3.4.1.2", adr_type="Riallocazione di sezione o clausola",
                            status="Revisionato DA", realloc_item=[c2.uuid], primary_proposed=d2.uuid,
                            proposer=t1.uuid, proposer_member="Anna Bianchi", da_decision="Commenti DA",
                            da_notes="Da rivedere.")
    decided = p.new_element("adr", title="Una scelta", adr_type="Altro / decisione trasversale", status="Decisa",
                            decision="Si fa così.")
    return p, dict(t1=t1, d1=d1, d2=d2, s1=s1, c1=c1, c2=c2, sf1=sf1, sf2=sf2, ei=ei, fe=fe, derived=derived,
                   chain=chain, realloc=realloc, decided=decided)


def test_the_templates_of_the_sa_phases_are_those_of_the_sheets():
    """Templates/ is up to date: written by scripts/make_sa_templates.py from the sheets of ATLAS."""
    for artifact in ARTIFACTS:
        path = ROOT / "Templates" / artifact.template
        assert path.exists(), f"{path.name}: run scripts/make_sa_templates.py"
        book = Workbook()
        write_sheet(book.active, Doc(filled=False, date="x"), artifact, None)
        expected, stored = values(book.active), values(openpyxl.load_workbook(path).active)
        del expected[1], stored[1]  # the line with the date of the template
        assert stored == expected, f"{path.name}: run scripts/make_sa_templates.py"


def test_the_template_of_a_reallocation_has_the_decision_of_the_da_in_its_section():
    path = next((ROOT / "Templates").glob("Template_X-1.06_*.xlsx"))
    labels = [row[0] for row in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]
    assert labels.index("Decisione della DA") < labels.index("Esito della decisione DA") < labels.index(
        "Osservazioni DA")
    assert "Descrizione" not in labels  # no description, options…: only the decision of the DA


def test_the_artefacts_are_filled_in_from_the_model(tmp_path):
    p, e = build(tmp_path)
    out = tmp_path / "out"
    files = generate(p, out, [a.code for a in ARTIFACTS], [e["realloc"], e["decided"]], True,
                     author="Salvatore De Simone", today=TODAY)
    assert [f.name for f in files] == [
        "Prova_SA_SA-1_Schede_di_dominio_funzionale_2026-10-30.xlsx",
        "Prova_SA_SA-2_Catalogo_funzioni_di_sistema_2026-10-30.xlsx",
        "Prova_SA_SA-3_Registro_tracciabilita_voci_SRS_2026-10-30.xlsx",
        "Prova_SA_SA-4_Requisiti_derivati_2026-10-30.xlsx",
        "Prova_SA_SA-5_Accordo_di_scambio_tra_domini_FE_2026-10-30.xlsx",
        "Prova_SA_SA-6_Dizionario_Exchange_Item_2026-10-30.xlsx",
        "Prova_SA_SA-7_Functional_chain_2026-10-30.xlsx",
        "Prova_SA_X-1_Decisioni_architetturali_ADR_2026-10-30.xlsx"]
    book = {f.name.split("_")[2]: openpyxl.load_workbook(f) for f in files}

    # SA-1: a sheet per domain, read again by «Importa clausole principali».
    sa1 = book["SA-1"]
    assert sa1.sheetnames == ["DF01", "DF02"]
    df01, df02 = sa1["DF01"], sa1["DF02"]
    assert df01["A1"].value == "SA-1 · Scheda di dominio funzionale · DF01"
    assert "Generato da ATLAS il 30/10/2026 dal progetto «Prova SA»" in df01["A2"].value
    assert [df01["B7"].value, df01["D7"].value, df01["B8"].value, df01["F8"].value] == [
        "SA-1", "T1 · DF01", "Salvatore De Simone", "30/10/2026"]
    assert find(df01, "Missione") == "Stimare la velocità."
    assert find(df01, "Referente") == "Anna Bianchi"
    assert find(df01, "Funzioni di sistema").startswith("SF-DF01-01 Stimare la velocità")
    assert find(df01, "Dati prodotti/posseduti dal dominio") == "EI-velocita Velocità stimata"
    assert find(df02, "Sezioni e clausole SRS a cui contribuisce") == "3.4.1.1 (DF01)"
    assert find(df02, "Dati usati da altri domini").startswith("EI-velocita Velocità stimata (DF01, FE-DF01-DF02-")
    assert [r["ID"] for r in table(df02, "Esito")] == ["ADR-001"]  # the domain proposed by the reallocation
    sheets = read_domain_sheets(files[0])
    assert [(s.ident, s.principal) for s in sheets] == [("DF01", "3.4.1"), ("DF02", "")]

    # SA-2: a sheet per domain, read again by «Importa funzioni»: nothing to change.
    sa2 = book["SA-2"]["DF01"]
    row = table(sa2, "ID")[0]
    assert [row["Sezioni SRS"], row["Righe 4.5.2"], row["Disponibilità nei modi (X / O)"],
            row["Voci di bordo tracciate"], row["Funzione nuova"], row["Motivo"]] == [
        "3.4.1", "4.5.2-R01", "FS: X, SR: O", "2", "Sì", "Manca nel catalogo"]
    from arcadia_tool.ui.import_functions import PRESENT, plan
    functions = read_function_sheets(files[1])
    assert [(f.sheet, f.name, f.clauses, f.origin, f.notes) for f in functions] == [
        ("DF01", "Stimare la velocità", "3.4.1", "Nuova", []), ("DF02", "Comandare i freni", "3.4.1.1", "", [])]
    assert [x.outcome for x in plan(p.model, functions)] == [PRESENT, PRESENT]

    # SA-3: one sheet, a row per item traced, with its section and domain, and its derived requirements.
    sa3 = table(book["SA-3"].active, "Sezione SRS")
    assert [(r["Sezione SRS"], r["Dominio principale"], r["ID voce"], r["Funzione primaria"],
             r["Funzioni secondarie"], r["Derivati / dizionario"]) for r in sa3] == [
        ("3.4.1", "DF01", "3.4.1.1", "SF-DF01-01", "SF-DF02-01", e["derived"].get("id"))]
    sa4 = table(book["SA-4"].active, "ID derivato")
    assert [(r["Voce madre"], r["Dominio"], r["Nome sintetico"]) for r in sa4] == [("3.4.1.1", "DF02", "Frenare")]
    sa5 = table(book["SA-5"].active, "ID scambio")
    assert [(r["Produttore (funzione, dominio)"], r["Utilizzatore (funzione, dominio)"], r["Chain che lo usano"])
            for r in sa5] == [("Stimare la velocità (DF01)", "Comandare i freni (DF02)", "FC-frenatura")]
    sa6 = table(book["SA-6"].active, "ID")
    assert sa6[0]["Produttore"] == "SF-DF01-01" and "V_TRAIN" in sa6[0][
        "Elementi (variabili cap. 7): variabile · descrizione · unità · campo · risoluzione"]
    chain = book["SA-7"]["FC-frenatura"]
    assert find(chain, "Domini coinvolti") == "DF01, DF02"
    assert [r["Funzione SF"] for r in table(chain, "#")] == ["SF-DF01-01", "SF-DF02-01"]
    assert table(chain, "Attraversamenti")[0]["Attraversamenti"] == "1"

    # The ADR: the register of the open ones, then a sheet per ADR with the decision of the DA.
    adrs = book["X-1"]
    assert adrs.sheetnames == ["ADR aperte", "ADR-001", "ADR-002"]
    assert [r["ID dell'ADR"] for r in table(adrs["ADR aperte"], "ID dell'ADR")] == ["ADR-001"]
    first = adrs["ADR-001"]
    assert find(first, "Sezioni o clausole da riallocare") == "3.4.1.2"
    assert find(first, "Dominio principale attuale") == "DF01 Odometria"  # from its section
    assert find(first, "Proponente (membro del team)") == "Anna Bianchi"
    assert find(first, "Esito della decisione DA") == "Commenti DA" and find(first, "Osservazioni DA") == "Da rivedere."
    assert find(adrs["ADR-002"], "Decisione") == "Si fa così."


def test_only_the_artefacts_with_data_and_of_the_domains_chosen(tmp_path):
    p, e = build(tmp_path)
    doc = make_doc(p, domains=[e["d2"].uuid], today=TODAY)
    assert [len(subjects(doc, a.code)) for a in ARTIFACTS] == [1, 1, 1, 1, 1, 1, 1]
    files = generate(p, tmp_path / "out", ["SA-1", "SA-2"], domains=[e["d2"].uuid], today=TODAY)
    assert openpyxl.load_workbook(files[0]).sheetnames == ["DF02"]
    empty = Project.create(tmp_path / "empty", "Vuoto")
    empty.new_element("domain", id="DF01", name="Solo")
    doc = make_doc(empty, today=TODAY)
    assert [len(subjects(doc, a.code)) for a in ARTIFACTS] == [1, 0, 0, 0, 0, 0, 0]
    assert generate(empty, tmp_path / "out2", ["SA-1", "SA-7"], today=TODAY)[0].name.startswith("Vuoto_SA-1")
    assert len(generate(empty, tmp_path / "out3", ["SA-7"], today=TODAY)) == 0  # nothing: no file


def test_the_settings_choose_what_to_generate_and_where(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QToolBar
    from arcadia_tool.documents.generate import load_settings
    from arcadia_tool.ui import documentation
    from arcadia_tool.ui.documentation import DocumentationDialog
    from arcadia_tool.ui.main_window import MainWindow

    p, e = build(tmp_path)
    p.new_element("domain", id="DF03", name="Senza nulla")
    window = MainWindow()
    window._set_project(p)
    assert window.sheets_action.isEnabled() and window.sheets_settings_action.isEnabled()
    assert window.sheets_action not in window.findChildren(QToolBar)[0].actions()  # in Progetto › Genera only
    # Without settings: «Schede Arcadia» says why, and offers them.
    told = []

    class Box(documentation.QMessageBox):
        def exec(self):
            told.append(self.text())
            return 0
    monkeypatch.setattr(documentation, "QMessageBox", Box)
    assert window.generate_sheets() is None and "non è ancora impostata" in told[0]
    shown = []
    monkeypatch.setattr(DocumentationDialog, "exec", lambda dialog: shown.append(dialog))
    window.sheets_settings()
    dialog = shown[0]
    assert dialog.sa_items["SA-3"].text(1) == "1 voce" and dialog.sa_items["SA-1"].text(1).startswith("3 domini")
    codes, adrs, register = dialog.selection()
    assert codes == [a.code for a in ARTIFACTS] and len(adrs) == 2 and register  # everything, at first
    # Only DF03, which has nothing but its sheet: the others cannot be chosen.
    for i in range(dialog.domains.count()):
        item = dialog.domains.item(i)
        item.setCheckState(Qt.Checked if item.text().startswith("DF03") else Qt.Unchecked)
    assert not dialog.all_domains.isChecked()
    assert dialog.sa_items["SA-2"].text(1) == "nessun dato nel modello"
    assert dialog.selection()[0] == ["SA-1"]
    assert dialog.sa_group.checkState(0) == Qt.Checked  # all those that can be chosen
    dialog.sa_group.setCheckState(0, Qt.Unchecked)  # the group: off, and back on
    assert dialog.selection()[0] == [] and dialog.sa_items["SA-2"].checkState(0) == Qt.Unchecked
    dialog.sa_group.setCheckState(0, Qt.Checked)
    assert dialog.selection()[0] == ["SA-1"] and dialog.sa_items["SA-2"].checkState(0) == Qt.Unchecked
    dialog.all_domains.setChecked(True)
    assert dialog.selection()[0] == [a.code for a in ARTIFACTS]  # available again: ticked
    dialog.sa_items["SA-4"].setCheckState(0, Qt.Unchecked)
    assert dialog.sa_group.checkState(0) == Qt.PartiallyChecked
    dialog.sa_items["SA-4"].setCheckState(0, Qt.Checked)
    dialog.sa_items["SA-3"].setCheckState(0, Qt.Unchecked)
    for item in dialog.adr_items.values():
        item.setCheckState(0, Qt.Unchecked)
    dialog.register_item.setCheckState(0, Qt.Unchecked)
    assert "X-1" not in dialog.summary.text() and "SA-3" not in dialog.summary.text()
    dialog.folder.setText(str(tmp_path / "docs"))
    assert dialog.save()
    saved = load_settings(p)
    assert saved["codes"] == ["SA-1", "SA-2", "SA-4", "SA-5", "SA-6", "SA-7"] and saved["adrs"] == []
    assert not saved["register"] and saved["domains"] is None and saved["folder"] == str(tmp_path / "docs")
    assert "documentation" in Project.open(p.path).settings  # with the project (project.yaml)
    # «Schede Arcadia»: as the settings say.
    written = documentation.generate_from_settings(window, p, ask=False)
    assert [Path(f).name.split("_")[2] for f in written] == ["SA-1", "SA-2", "SA-4", "SA-5", "SA-6", "SA-7"]
    asked = []
    monkeypatch.setattr(Box, "question", staticmethod(lambda *a, **k: asked.append(a[2]) or Box.No))
    assert window.generate_sheets() is None and "SA-1" in asked[0]  # already there: asked, not replaced
    # Opened again: the choice saved; all the ADR, also one created later.
    again = DocumentationDialog(p)
    assert again.folder.text() == str(tmp_path / "docs")
    assert again.sa_items["SA-3"].checkState(0) == Qt.Unchecked and again.sa_items["SA-1"].checkState(0) == Qt.Checked
    for item in again.adr_items.values():
        item.setCheckState(0, Qt.Checked)
    again.save()
    assert load_settings(p)["adrs"] == "all"
    later = p.new_element("adr", title="Dopo", status="Proposta")
    from arcadia_tool.documents.generate import chosen
    assert later in chosen(p, load_settings(p))[1]
    window._set_project(None)
    window.close()


def test_an_artefact_without_data_stays_wanted_until_it_has_some(app, tmp_path):
    from arcadia_tool.documents.generate import load_settings, problems
    from arcadia_tool.ui.documentation import DocumentationDialog

    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Solo")
    dialog = DocumentationDialog(p)
    dialog.folder.setText(str(tmp_path / "docs"))
    assert dialog.sa_items["SA-2"].text(1) == "nessun dato nel modello"
    dialog.save()
    assert load_settings(p)["codes"] == [a.code for a in ARTIFACTS]  # wanted: generated when they have data
    assert not problems(p, load_settings(p))
    dialog = DocumentationDialog(p)
    dialog.sa_items["SA-1"].setCheckState(0, Qt.Unchecked)
    for item in dialog.adr_items.values():
        item.setCheckState(0, Qt.Unchecked)
    dialog.register_item.setCheckState(0, Qt.Unchecked)
    dialog.save()
    assert problems(p, load_settings(p)) == ["Nessuno degli artefatti scelti ha dati nel modello."]
    p.new_element("system_function", name="Una funzione", domain=d1.uuid)
    assert not problems(p, load_settings(p))  # SA-2 has data now
