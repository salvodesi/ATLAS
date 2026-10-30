"""The derived requirements (SA-4) mapped on a system function (SA-2): «Mappa requisiti derivati…» of its sheet, a
window with only the clauses they are derived from and, under each, its derived requirements with ticks."""
import pytest
from PySide6.QtWidgets import QApplication, QPushButton

from arcadia_tool.storage.project import Project


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def build(tmp_path):
    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Protezione")
    d2 = p.new_element("domain", id="DF02", name="Localizzazione")
    s1 = p.new_element("requirement", id="3.4.1", item_type="Sezione", sequence=1)
    c1 = p.new_element("requirement", id="3.4.1.1", item_type="Requisito", parent=s1.uuid, sequence=2,
                       category=["Bordo"])
    bullet = p.new_element("requirement", id="3.4.1.1 a)", item_type="Voce", parent=c1.uuid, sequence=3)
    c2 = p.new_element("requirement", id="3.4.1.2", item_type="Requisito", parent=s1.uuid, sequence=4)
    s5 = p.new_element("requirement", id="5.1.1", item_type="Sezione", sequence=10)
    c5 = p.new_element("requirement", id="5.1.1.1", item_type="Requisito", parent=s5.uuid, sequence=11)
    sf1 = p.new_element("system_function", name="Proteggere", domain=d1.uuid)
    sf2 = p.new_element("system_function", name="Localizzare", domain=d2.uuid)
    derived = {
        "own": p.new_element("derived_requirement", parent=c1.uuid, obligation="Frena.", domain=d1.uuid),
        "bullet": p.new_element("derived_requirement", parent=bullet.uuid, obligation="Avvisa.", domain=d1.uuid),
        "taken": p.new_element("derived_requirement", parent=c5.uuid, obligation="Calcola.",
                               system_function=sf2.uuid),
        "other": p.new_element("derived_requirement", parent=c5.uuid, obligation="Riporta.", domain=d2.uuid),
    }
    return p, dict(d1=d1, d2=d2, s1=s1, c1=c1, c2=c2, c5=c5, sf1=sf1, sf2=sf2, **derived)


def test_the_window_lists_only_the_clauses_with_derived_requirements(app, tmp_path, monkeypatch):
    import arcadia_tool.ui.srs_trace as srs_trace

    p, e = build(tmp_path)
    assert e["taken"].get("domain") == e["d2"].uuid  # from its system function
    dialog = srs_trace.TraceDerivedDialog(p, e["sf1"])
    listed = {x.uuid for items in dialog.by_chapter.values() for x in items}
    assert {e["c1"].uuid, e["c5"].uuid} <= listed  # the clauses they are derived from (a bullet: its clause)
    assert not {e["s1"].uuid, e["c2"].uuid} & listed  # no other clause, no section
    assert set(dialog.state) == {e[k].uuid for k in ("own", "bullet", "taken", "other")}  # only these are ticked
    from arcadia_tool.ui.srs_reader import chapter_of
    dialog.tabs.setCurrentIndex(next(i for i in range(dialog.tabs.count())
                                     if dialog.tabs.tabData(i) == chapter_of(e["c1"])))
    row = dialog._items[e["c1"].uuid]
    assert {row.child(i).data(0, srs_trace.UUID_ROLE) for i in range(row.childCount())} == {
        e["own"].uuid, e["bullet"].uuid}  # under their clause
    assert row.icon(2).isNull()  # a clause: nothing to tick
    for key in ("own", "taken", "other"):
        dialog.set_state(e[key].uuid, srs_trace.ALLOCATED)
    assert dialog.count.text() == "3 requisiti derivati mappati"
    asked = []
    monkeypatch.setattr(srs_trace.QMessageBox, "question",
                        lambda *a, **k: asked.append(a[1]) or srs_trace.QMessageBox.Yes)
    dialog.accept()
    # Asked: one of another domain (it takes the domain of the function), one mapped on another function.
    assert asked == ["Requisiti derivati di un altro dominio", "Requisito derivato già mappato"]
    assert [e[k].get("system_function") for k in ("own", "bullet", "taken", "other")] == [
        e["sf1"].uuid, None, e["sf1"].uuid, e["sf1"].uuid]
    assert e["other"].get("domain") == e["d1"].uuid and e["taken"].get("domain") == e["d1"].uuid
    assert e["other"] in dialog.changed and e["own"].uuid in p.pending()
    # Unticked: no longer mapped.
    again = srs_trace.TraceDerivedDialog(p, e["sf1"])
    assert again.chosen(srs_trace.ALLOCATED) == {e[k].uuid for k in ("own", "taken", "other")}
    again.set_state(e["own"].uuid, None)
    again.accept()
    assert e["own"].get("system_function") is None and e["taken"].get("system_function") == e["sf1"].uuid
    # Cancelled: nothing changes.
    cancelled = srs_trace.TraceDerivedDialog(p, e["sf1"])
    cancelled.set_state(e["taken"].uuid, None)
    cancelled.reject()
    assert e["taken"].get("system_function") == e["sf1"].uuid


def test_the_sheets_show_the_mapping_and_the_coverage_counts_it(app, tmp_path, monkeypatch):
    import arcadia_tool.ui.srs_trace as srs_trace
    from arcadia_tool.model.coverage import FUNCTION, coverage
    from arcadia_tool.ui.main_window import MainWindow
    from arcadia_tool.ui.srs_reader import document_order

    p, e = build(tmp_path)
    window = MainWindow()
    window._set_project(p)
    window._select(e["sf1"])
    editor = window.editor
    buttons = {b.text(): b for b in editor.widget().findChildren(QPushButton)}
    assert "Mappa requisiti derivati…" in buttons and "Alloca a sezioni e clausole SRS…" in buttons

    def ticks(dialog):  # the window: «own» ticked, then «Applica»
        dialog.set_state(e["own"].uuid, srs_trace.ALLOCATED)
        dialog.accept()
        return True
    monkeypatch.setattr(srs_trace.TraceDerivedDialog, "exec", ticks)
    buttons["Mappa requisiti derivati…"].click()
    assert e["own"].get("system_function") == e["sf1"].uuid
    from PySide6.QtWidgets import QListWidget
    mapped = window.editor.widget().findChild(QListWidget, "derivedMapped")
    assert [mapped.item(i).data(0x0100) for i in range(mapped.count())] == [e["own"].uuid]
    window.save_tab()  # Ctrl+S on the sheet: the derived requirement mapped is written
    assert e["own"].uuid not in p.pending()
    # The sheet of the derived requirement: its system function, chosen with ticks.
    window._select(e["own"])
    assert window.editor.widgets["systemFunction"].text() == p.model.choice_label(e["sf1"])
    # The item it is derived from: covered at the level of the system functions too.
    rows = {r.item.get("id"): r for r in coverage(p.model, document_order(p.model), "Bordo").rows}
    assert rows["3.4.1.1"].at(FUNCTION) and e["sf1"] in rows["3.4.1.1"].functions
    window._set_project(None)
    window.close()


def test_the_button_is_there_without_derived_requirements_and_their_list_is_under(app, tmp_path):
    import arcadia_tool.ui.srs_trace as srs_trace
    from PySide6.QtWidgets import QListWidget
    from arcadia_tool.ui.main_window import MainWindow

    p = Project.create(tmp_path / "p", "Prova")
    d1 = p.new_element("domain", id="DF01", name="Protezione")
    sf = p.new_element("system_function", name="Proteggere", domain=d1.uuid)
    window = MainWindow()
    window._set_project(p)
    window._select(sf)
    page = window.editor.widget()
    buttons = {b.text(): b for b in page.findChildren(QPushButton)}
    assert buttons["Mappa requisiti derivati…"].isEnabled()  # also with no derived requirement in the model
    dialog = srs_trace.TraceDerivedDialog(p, sf)
    assert not dialog.state  # nothing to map yet
    window.resize(1400, 1600)
    window.show()
    app.processEvents()
    derived = page.findChild(QListWidget, "derivedMapped")
    srs = page.findChild(QListWidget, "srsAllocated")
    top = lambda view: view.mapTo(page, view.rect().topLeft())
    assert top(srs).y() + srs.height() <= top(derived).y()  # one under the other, not side by side
    window._set_project(None)
    window.close()


def test_the_documents_have_the_mapping(tmp_path):
    import openpyxl
    from arcadia_tool.documents.generate import generate

    p, e = build(tmp_path)
    e["own"].set("system_function", e["sf1"].uuid)
    p.update(e["own"])
    files = generate(p, tmp_path / "out", ["SA-2", "SA-4"])
    sa2 = openpyxl.load_workbook(files[0])["DF01"]
    header = next(r for r in sa2.iter_rows(values_only=True) if r[0] == "ID")
    row = next(r for r in sa2.iter_rows(values_only=True) if r[1] == "Proteggere")
    assert row[header.index("Requisiti derivati (SA-4)")] == e["own"].get("id")
    sa4 = openpyxl.load_workbook(files[1]).active
    header = next(r for r in sa4.iter_rows(values_only=True) if r[0] == "ID derivato")
    row = next(r for r in sa4.iter_rows(values_only=True) if r[0] == e["own"].get("id"))
    assert row[header.index("Funzione di sistema")] == e["sf1"].get("id")
