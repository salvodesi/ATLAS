"""The list of the open architectural decisions (X-1) as an Excel file: those not yet decided by the Design
Authority, with the team that opened them, in the style of the templates of the method."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from arcadia_tool.model.element import natural_key

CLOSED = ("Decisa", "Superata")
COLUMNS = (("ID", 14), ("Descrizione", 70), ("Team proponente", 34), ("Stato", 16))
DARK, LIGHT, LINE, MUTED = "1F3864", "D9E2F3", "8EA9DB", "7F7F7F"


def open_adrs(model):
    """The ADR still open (proposed, or without a status), in the order of their IDs: (id, title, team, status)."""
    rows = []
    for adr in model.of_kind("adr"):
        status = adr.get("status") or ""
        if status in CLOSED:
            continue
        team = model.get(adr.get("proposer")) if adr.get("proposer") else None
        proposer = " · ".join(t for t in (model.label(team) if team is not None else "",
                                          adr.get("proposer_member") or "") if t)  # the team, its member
        rows.append((adr.get("id") or "", adr.get("title") or "", proposer, status or "Proposta"))
    return sorted(rows, key=lambda r: natural_key(r[0]))


def write_open_adrs(model, path, project_name="", today=None):
    """Write the list of the open ADR to the Excel file `path`; returns how many they are."""
    rows = open_adrs(model)
    today = today or date.today()
    book = Workbook()
    ws = book.active
    ws.title = "ADR aperte"
    side = Side(style="thin", color=LINE)
    border = Border(left=side, right=side, top=side, bottom=side)
    ws["A1"] = "X-1 · Decisioni architetturali (ADR) aperte"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = (f"Progetto {project_name} · " if project_name else "") + \
        f"Estratto da ATLAS il {today:%d/%m/%Y} · Le ADR non ancora decise dalla Design Authority: {len(rows)}"
    ws["A2"].font = Font(name="Arial", size=9, color=MUTED)
    for c, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=4, column=c, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    for r, row in enumerate(rows, start=5):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Arial", size=10, bold=c == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if c == 1:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    if rows:  # filters and sorting on the columns
        table = Table(displayName="ADR_aperte", ref=f"A4:D{4 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(table)
    else:
        ws["A5"] = "Nessuna ADR aperta."
        ws["A5"].font = Font(name="Arial", size=10, italic=True, color=MUTED)
    ws.freeze_panes = "A5"
    ws.row_dimensions[4].height = 20
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    book.properties.title = "X-1 · ADR aperte"
    book.save(path)
    return len(rows)
