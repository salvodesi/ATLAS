"""The coverage report of the SRS items (Bordo, Definizione) as an Excel file: a summary sheet, then for each
category the coverage of each section and every clause (with its chapter 9 category, the section containing it,
the domains, system functions and derived requirements that cover it)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from arcadia_tool.model.coverage import DERIVED, DOMAIN, FUNCTION, LEVELS, subset_coverage

DARK, LIGHT, LINE, MUTED = "1F3A5F", "EEF3F9", "C9D1DC", "6B7280"
GREEN, RED = "166534", "B42318"
COLUMNS = (("Voce SRS", 14), ("Sezione", 11), ("Categoria", 16), ("Testo", 70), ("Dominio funzionale", 34),
           ("Funzioni di sistema", 34), ("Requisiti derivati", 26), ("Copertura", 12))
SECTION_COLUMNS = (("Sezione", 11), ("Titolo", 50), ("Clausole", 10), ("Coperte · dominio", 18),
                   ("Coperte · funzione di sistema", 18), ("Coperte · requisito derivato", 18),
                   ("Scoperte (dominio)", 12))


def _header(sheet, columns, border):
    for c, (label, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=c, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[cell.column_letter].width = width


def _table(sheet, name, columns, rows):
    if rows:
        last = sheet.cell(row=1, column=len(columns)).column_letter
        table = Table(displayName=name, ref=f"A1:{last}{1 + rows}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        sheet.add_table(table)
    sheet.freeze_panes = "B2"


def share(report, level):
    return f"{report.count(level)} ({report.percent(level):.1f}%)"


def domains_text(model, row):
    text = ", ".join(f"{model.label(d)}{' (principale)' if principal else ''}" for d, principal in row.domains)
    if text and row.domains_from is not None:
        text += f" (da {row.domains_from.get('id')})"
    return text


def write_coverage(model, reports, path, project_name="", today=None):
    """Write the report (the Coverage of each category) to the Excel file path."""
    today = today or date.today()
    side = Side(style="thin", color=LINE)
    border = Border(left=side, right=side, top=side, bottom=side)
    book = Workbook()
    ws = book.active
    ws.title = "Riepilogo"
    ws["A1"] = "Copertura delle voci SRS"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = (f"Progetto {project_name} · " if project_name else "") + f"Estratto da ATLAS il {today:%d/%m/%Y}"
    ws["A2"].font = Font(name="Arial", size=9, color=MUTED)
    covered, total, percent = subset_coverage(reports)
    ws["A4"] = "Copertura del subset (a livello di dominio)"
    ws["A4"].font = Font(name="Arial", size=11, bold=True)
    ws["B4"] = f"{percent:.1f}% ({covered} di {total} voci)"
    ws["B4"].font = Font(name="Arial", size=11, bold=True, color=GREEN if covered == total else DARK)
    headers = ["Categoria", "Voci"] + [f"Coperte: {label}" for _level, label in LEVELS] + ["Scoperte (dominio)"]
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=c, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 26 if c > 1 else 18
    for r, report in enumerate(reports, start=7):
        values = [report.category, report.total] + \
            [f"{report.count(level)} ({report.percent(level):.1f}%)" for level, _label in LEVELS] + \
            [len(report.uncovered)]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Arial", size=10, bold=c == 1)
            cell.border = border
    for report in reports:
        sections = book.create_sheet(f"{report.category} - sezioni")
        _header(sections, SECTION_COLUMNS, border)
        groups = report.sections()
        for r, s in enumerate(groups, start=2):
            values = (s.section.get("id") if s.section is not None else "(fuori sezione)",
                      (s.section.get("title") or "") if s.section is not None else "", s.total,
                      share(s, DOMAIN), share(s, FUNCTION), share(s, DERIVED), len(s.uncovered))
            for c, value in enumerate(values, start=1):
                cell = sections.cell(row=r, column=c, value=value)
                cell.font = Font(name="Arial", size=10, bold=c in (1, 4), color=(
                    (GREEN if s.count(DOMAIN) == s.total else RED) if c == 4 else
                    RED if c == len(values) and value else None))
                cell.border = border
        _table(sections, f"Sezioni_{report.category}", SECTION_COLUMNS, len(groups))
        sheet = book.create_sheet(f"{report.category} - clausole")
        _header(sheet, COLUMNS, border)
        for r, row in enumerate(report.rows, start=2):
            values = (row.item.get("id"), row.section.get("id") if row.section is not None else "",
                      ", ".join(row.item.get("category") or []),
                      " ".join((row.item.get("text") or row.item.get("title") or "").split()),
                      domains_text(model, row), ", ".join(model.label(f) for f in row.functions),
                      ", ".join(model.label(d) for d in row.derived), "Coperta" if row.covered else "Scoperta")
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=c, value=value)
                cell.font = Font(name="Arial", size=10, bold=c == 1,
                                 color=(GREEN if row.covered else RED) if c == len(values) else None)
                cell.alignment = Alignment(vertical="top", wrap_text=c == 4)
                cell.border = border
                if c == 1:
                    cell.fill = PatternFill("solid", fgColor=LIGHT)
        _table(sheet, f"Clausole_{report.category}", COLUMNS, len(report.rows))
    book.save(path)
    return path

