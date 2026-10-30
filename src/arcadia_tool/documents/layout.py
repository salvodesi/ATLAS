"""The layout of a sheet of the method, in the style of the templates ARCADIA: the title and the guide (rows 1–5),
the data of the document (rows 7–9), then sections, fields and tables, and the checklist. The same layout writes a
template (empty values: the guide in grey italics) and a document filled in from the model."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FIRST_ROW = 11  # the first row after the data of the document
DARK, LIGHT, LINE, GREY = "1F3864", "D9E2F3", "8EA9DB", "7F7F7F"
LEGEND = ("Legenda: le celle bianche vanno compilate; il testo grigio in corsivo è guida o esempio del template e va "
          "sostituito. Le celle «Stato» e le colonne a valori fissi propongono un elenco a tendina.")
_SIDE = Side(style="thin", color=LINE)
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
_WRAP = Alignment(vertical="center", wrap_text=True)


def _font(size=10, bold=False, italic=False, color=None):
    return Font(name="Arial", size=size, bold=bold, italic=italic, color=color)


STYLES = {
    "label": (_font(bold=True), PatternFill("solid", fgColor=LIGHT), _WRAP),
    "value": (_font(), None, _WRAP),
    "hint": (_font(italic=True, color=GREY), None, _WRAP),
    "header": (_font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=DARK),
               Alignment(horizontal="center", vertical="center", wrap_text=True)),
    "cell": (_font(), None, Alignment(horizontal="left", vertical="top", wrap_text=True)),
    "cell_hint": (_font(italic=True, color=GREY), None, Alignment(horizontal="left", vertical="top", wrap_text=True)),
}


class SheetWriter:
    """Writes the rows of a sheet one after the other (self.row: the next one). widths: of the columns, A first."""

    def __init__(self, ws, widths):
        self.ws, self.widths = ws, list(widths)
        self.n = len(self.widths)
        self.last = get_column_letter(self.n)
        for c, width in enumerate(self.widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.sheet_view.showGridLines = False
        self.row = 1

    # --- cells -----------------------------------------------------------------------------------------------

    def _style(self, cell, style, border=True):
        font, fill, alignment = STYLES[style]
        cell.font, cell.alignment = font, alignment
        if fill is not None:
            cell.fill = fill
        if border:
            cell.border = _BORDER

    def _span(self, first, last, r, value, style):
        """Columns first..last (1-based) of row r as one cell."""
        for c in range(first, last + 1):
            self._style(self.ws.cell(r, c), style)
        self.ws.cell(r, first).value = value if value not in ("", None) else None
        if last > first:
            self.ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)

    def choose(self, ref, options):
        """A drop-down list of the values of the cell."""
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        validation.add(ref)
        self.ws.add_data_validation(validation)

    @staticmethod
    def _height(texts, widths):
        """Rows of text the longest cell needs (about one character per unit of width)."""
        lines = 1
        for text, width in zip(texts, widths):
            for part in str(text or "").split("\n"):
                lines = max(lines, -(-len(part) // max(int(width * 1.1), 8)))
            lines = max(lines, str(text or "").count("\n") + 1)
        return min(15 + 13 * lines, 400)

    # --- the parts of a sheet ---------------------------------------------------------------------------------

    def intro(self, title, line, purpose, guide):
        """Rows 1–5: the title, the line of the template (version, baseline, who fills it in), the purpose, how it is
        filled in, the legend."""
        ws = self.ws
        texts = (title, line, f"Scopo. {purpose}", f"Come si compila. {guide}", LEGEND)
        fonts = (_font(14, bold=True), _font(9, color=GREY), _font(9), _font(9), _font(9, italic=True, color=GREY))
        for r, (text, font) in enumerate(zip(texts, fonts), 1):
            cell = ws.cell(r, 1, text)
            cell.font, cell.alignment = font, Alignment(vertical="top", wrap_text=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=self.n)
            ws.row_dimensions[r].height = 22 if r == 1 else self._height([text], [sum(self.widths)]) - 2
        self.row = 7

    def document(self, values):
        """Rows 7–9: the data of the document, three label and value pairs per row. values: {label: (text,
        is_hint)}."""
        rows = (("Template", "Team / dominio", "Versione"), ("Autore", "Revisore", "Data"),
                ("Stato", "Baseline SRS", "Rif. Capella / DOORS"))
        middle = max((self.n - 4) // 2, 1)  # the value of the second pair: from D
        spans = ((1, 2, 2), (3, 4, 3 + middle), (4 + middle, 5 + middle, self.n))
        for r, labels in enumerate(rows, 7):
            texts, widths = [], []
            for label, (at, first, last) in zip(labels, spans):
                text, is_hint = values.get(label, ("", False))
                self._span(at, at, r, label, "label")
                self._span(first, max(first, last), r, text, "hint" if is_hint else "value")
                if label == "Stato" and is_hint:
                    self.choose(f"{get_column_letter(first)}{r}", ("Bozza", "In revisione", "Approvato"))
                texts += [label, text]
                widths += [self.widths[at - 1] * 0.85, sum(self.widths[first - 1:max(first, last)])]
            self.ws.row_dimensions[r].height = max(18, self._height(texts, widths) - 10)
        self.row = FIRST_ROW

    def heading(self, text):
        """A dark row across the sheet: the start of a section."""
        if self.row > FIRST_ROW:
            self.row += 1
        self._span(1, self.n, self.row, text, "header")
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def field(self, label, text="", hint="", options=(), tall=False):
        """A label in A and its value across the other columns: the text, else the guide in grey italics."""
        r = self.row
        self._span(1, 1, r, label, "label")
        self._span(2, self.n, r, text or hint, "value" if text else "hint")
        if options and not text:
            self.choose(f"B{r}", options)
        height = self._height([label, text or hint], [self.widths[0], sum(self.widths[1:])])
        self.ws.row_dimensions[r].height = max(height, 45 if tall and not text else 22)
        self.row += 1

    def table(self, columns, rows=None, example=None, empty=5, label=None):
        """A table: its columns [(heading, number of sheet columns, options)], then its rows of texts; without rows
        (a template), the example in grey italics and empty rows to fill in. label: a caption above it."""
        if label:
            self._span(1, self.n, self.row, label, "label")
            self.ws.row_dimensions[self.row].height = 22
            self.row += 1
        spans, c = [], 1
        for i, (_heading, size, _options) in enumerate(columns):
            last = self.n if i == len(columns) - 1 else c + size - 1
            spans.append((c, last))
            c = last + 1
        widths = [sum(self.widths[a - 1:b]) for a, b in spans]
        headings = [h for h, _s, _o in columns]
        self._row(headings, spans, "header", max(30, self._height(headings, [w * 0.85 for w in widths])))
        filled = rows is not None
        body = rows if filled else ([example] if example else []) + [[""] * len(columns)] * empty
        for i, values in enumerate(body):
            style = "cell_hint" if not filled and example and i == 0 else "cell"
            r = self._row(values, spans, style, self._height(values, widths))
            if not filled:
                for (first, _last), (_h, _s, options) in zip(spans, columns):
                    if options:
                        self.choose(f"{get_column_letter(first)}{r}", options)
        if filled and not rows:
            self._row(["—"] + [""] * (len(columns) - 1), spans, "cell_hint", 20)

    def _row(self, values, spans, style, height):
        r = self.row
        for (first, last), value in zip(spans, values):
            self._span(first, last, r, value, style)
        self.ws.row_dimensions[r].height = height
        self.row += 1
        return r

    def checklist(self, title, items, ticked=()):
        """The checklist of the sheet: ☐ and the check, across the other columns."""
        self.row += 1
        cell = self.ws.cell(self.row, 1, title)
        cell.font, cell.alignment = _font(11, bold=True), Alignment(vertical="top", wrap_text=True)
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.n)
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1
        for item in items:
            box = self.ws.cell(self.row, 1, "☑" if item in ticked else "☐")
            self._style(box, "value")
            box.font, box.alignment = _font(12), Alignment(horizontal="center", vertical="center")
            self._span(2, self.n, self.row, item, "value")
            self.ws.row_dimensions[self.row].height = self._height([item], [sum(self.widths[1:])])
            self.row += 1

    def finish(self):
        """Printed on landscape pages as wide as the sheet."""
        ws = self.ws
        ws.print_area = f"A1:{self.last}{self.row - 1}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
