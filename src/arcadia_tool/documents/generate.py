"""«Genera documentazione»: the artefacts of the SA phases and the ADR chosen, filled in from the model, written as
Excel files in a folder (one file per artefact; a sheet per domain, chain or ADR, or one sheet for the whole
table)."""
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from arcadia_tool.documents.adr import WHO, write_adr, write_register
from arcadia_tool.documents.sa import ARTIFACTS, ARTIFACTS_BY_CODE, BASELINE, Doc, write_sheet
from arcadia_tool.documents.values import by_id

ADR_FILE = "X-1_Decisioni_architetturali_ADR"
REGISTER = "register"  # in the choice of the ADR: the register of the open ADR


def make_doc(project, author="", domains=None, today=None):
    today = today or date.today()
    return Doc(model=project.model, project=project.name, baseline=project.baseline or BASELINE, author=author or "",
               date=f"{today:%d/%m/%Y}", domains=set(domains) if domains is not None else None)


def subjects(doc, code):
    """What the artefact has in the model (with the domains chosen): its sheets or its rows; empty: nothing."""
    return list(ARTIFACTS_BY_CODE[code].subjects(doc))


def _ascii(text):
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return "_".join("".join(c if c.isalnum() or c in "-_" else " " for c in text).split()) or "Progetto"


def file_name(project_name, stem, today=None):
    """«<progetto>_<artefatto>_<data>.xlsx»."""
    return f"{_ascii(project_name)}_{stem}_{(today or date.today()):%Y-%m-%d}.xlsx"


def artifact_file(project_name, code, today=None):
    artifact = ARTIFACTS_BY_CODE[code]
    return file_name(project_name, f"{code}_{artifact.files}", today)


def adr_file(project_name, today=None):
    return file_name(project_name, ADR_FILE, today)


def _sheet_name(text, used):
    """A name valid for a worksheet (at most 31 characters, without : \\ / ? * [ ]), not used yet in the book."""
    base = "".join("-" if c in ':\\/?*[]' else c for c in str(text)).strip()[:31] or "Foglio"
    name, n = base, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(name.lower())
    return name


def write_artifact(doc, code, path):
    """The file of an SA artefact; None if the model has nothing for it."""
    artifact = ARTIFACTS_BY_CODE[code]
    found = subjects(doc, code)
    if not found:
        return None
    book = Workbook()
    book.remove(book.active)
    used = set()
    if artifact.per:
        for subject in found:
            ws = book.create_sheet(_sheet_name(subject.get("id") or subject.get("title") or code, used))
            write_sheet(ws, doc, artifact, subject)
    else:
        write_sheet(book.create_sheet(_sheet_name(code, used)), doc, artifact, found)
    book.properties.title = f"{code} · {artifact.title} · {doc.project}"
    book.save(path)
    return path


def write_adrs(doc, adrs, register, path):
    """The file of the ADR: the register of the open ones first (register), then a sheet per ADR chosen."""
    model = doc.model
    if not adrs and not register:
        return None
    book = Workbook()
    book.remove(book.active)
    used = set()
    line = (f"Generato da ATLAS il {doc.date} dal progetto «{doc.project}» · Baseline delle specifiche: "
            f"{doc.baseline} · {WHO}")
    if register:
        write_register(book.create_sheet(_sheet_name("ADR aperte", used)), doc, line)
    for adr in by_id(adrs):
        write_adr(book.create_sheet(_sheet_name(adr.get("id") or "ADR", used)), doc, adr.get("adr_type"), adr, line)
    book.properties.title = f"X-1 · Decisioni architetturali · {doc.project}"
    book.save(path)
    return path


# --- the settings of «Genera › Schede Arcadia», saved with the project -------------------------------------------
# What to generate is shared with the team (project.yaml, «documentation»); the folder is of this PC (.arcadia/).
SETTINGS_KEY, FOLDER_KEY, ALL = "documentation", "documentation_folder", "all"


def load_settings(project):
    """{codes, adrs ("all" or uuids), register, domains (None: all), folder}; None if never saved."""
    data = project.settings.get(SETTINGS_KEY)
    if not isinstance(data, dict):
        return None
    return {"codes": [c for c in data.get("codes") or [] if c in ARTIFACTS_BY_CODE],
            "adrs": data.get("adrs") if data.get("adrs") == ALL else list(data.get("adrs") or []),
            "register": bool(data.get("register")), "domains": data.get("domains"),
            "folder": project.local_settings().get(FOLDER_KEY, "")}


def save_settings(project, codes, adrs, register, domains, folder):
    """adrs: "all" (also those created later) or the uuids of some; domains: None (all) or uuids."""
    project.update_settings(**{SETTINGS_KEY: {"codes": [a.code for a in ARTIFACTS if a.code in codes],
                                              "adrs": adrs if adrs == ALL else list(adrs), "register": bool(register),
                                              "domains": list(domains) if domains is not None else None}})
    project.set_local_settings(**{FOLDER_KEY: str(folder or "")})


def chosen(project, settings):
    """(codes, ADR, register, domains, folder) of the settings, as the model is now: the ADR deleted left out, the
    domains no longer there too."""
    model = project.model
    adrs = model.of_kind("adr") if settings["adrs"] == ALL else [model.get(u) for u in settings["adrs"]]
    domains = settings["domains"]
    if domains is not None:
        domains = [u for u in domains if model.get(u) is not None]
    return settings["codes"], [a for a in adrs if a is not None], settings["register"], domains, settings["folder"]


def problems(project, settings, author=""):
    """Why «Schede Arcadia» cannot generate with these settings (empty: it can)."""
    if settings is None:
        return ["La documentazione da generare non è ancora impostata (Progetto › Genera › Impostazioni…)."]
    codes, adrs, register, domains, folder = chosen(project, settings)
    found = []
    if not str(folder or "").strip():
        found.append("Manca la cartella di destinazione.")
    doc = make_doc(project, author, domains)
    with_data = [c for c in codes if subjects(doc, c)]
    if not with_data and not adrs and not (register and project.model.of_kind("adr")):
        found.append("Nessuno degli artefatti scelti ha dati nel modello." if codes or register else
                     "Non è scelto nessun artefatto da generare.")
    return found


def planned_files(project, folder, codes, adrs, register, today=None):
    """The files «Genera» writes in the folder."""
    folder = Path(folder)
    files = [folder / artifact_file(project.name, code, today) for code in codes]
    if adrs or register:
        files.append(folder / adr_file(project.name, today))
    return files


def generate(project, folder, codes, adrs=(), register=False, author="", domains=None, progress=None, today=None):
    """Writes the artefacts chosen (codes: SA-1…SA-7; adrs: the ADR; register: the open ADR) in the folder.
    progress(done, total, what) after each file. Returns the paths written (an artefact without data is skipped)."""
    doc = make_doc(project, author, domains, today)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    order = [a.code for a in ARTIFACTS if a.code in codes]
    total = len(order) + (1 if adrs or register else 0)
    written = []
    for n, code in enumerate(order):
        if progress:
            progress(n, total, f"{code} · {ARTIFACTS_BY_CODE[code].title}")
        path = write_artifact(doc, code, folder / artifact_file(project.name, code, today))
        if path is not None:
            written.append(path)
    if adrs or register:
        if progress:
            progress(len(order), total, "X-1 · Decisioni architetturali (ADR)")
        written.append(write_adrs(doc, list(adrs), register, folder / adr_file(project.name, today)))
    if progress:
        progress(total, total, "Fatto")
    return written
